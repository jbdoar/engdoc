#!/usr/bin/env python3
"""Validate an engdoc project set and export its requirements traceability."""

import argparse
from datetime import date
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DOCUMENTS = Path(__file__).resolve().parent
sys.path[:0] = [
    str(DOCUMENTS), str(DOCUMENTS / "requirements"),
    str(DOCUMENTS / "calibration"), str(DOCUMENTS / "risk"),
]

from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata, read_metadata
from engdoc_org import read_headings
from export_requirements import parse_requirements, validate_requirements
from export_calibration import parse_equipment, validate_equipment
from export_risk import parse_hazards, validate_hazards


REFERENCE_PROPERTY = "REQUIREMENTS"
REFERENCE_SPLIT = re.compile(r"[\s,]+")


def split_references(value):
    """Return stable, nonempty identifiers from a reference property VALUE."""
    return list(dict.fromkeys(item for item in REFERENCE_SPLIT.split(value.strip()) if item))


def load_project_set(directory):
    """Load controlled Org documents directly beneath DIRECTORY."""
    documents = []
    for path in sorted(Path(directory).glob("*.org")):
        metadata = read_metadata(path)
        if not metadata.get("ENGDOC_DOCUMENT_ID"):
            continue
        metadata, warnings = load_and_validate_metadata(path)
        documents.append({"path": path, "metadata": metadata, "warnings": warnings})
    if not documents:
        raise ValueError("No controlled engdoc documents found")
    return documents


def validate_project_set(directory):
    """Validate document and requirement links; return an export model."""
    documents = load_project_set(directory)
    by_id = {}
    project_ids = set()
    warnings = []

    for document in documents:
        metadata = document["metadata"]
        document_id = metadata["ENGDOC_DOCUMENT_ID"]
        if document_id in by_id:
            raise ValueError(f"Duplicate document ID: {document_id}")
        by_id[document_id] = document
        project_ids.add(metadata["ENGDOC_PROJECT_ID"])
        warnings.extend(
            f"{document_id}: {warning}" for warning in document["warnings"]
        )

    if len(project_ids) != 1:
        raise ValueError(
            "Documents use different project IDs: " + ", ".join(sorted(project_ids))
        )

    for document_id, document in by_id.items():
        related = split_references(
            document["metadata"].get("ENGDOC_RELATED_DOCUMENTS", "")
        )
        for target in related:
            if target == document_id:
                raise ValueError(f"{document_id}: related document cannot reference itself")
            if target not in by_id:
                raise ValueError(f"{document_id}: unknown related document {target}")
        document["related"] = related

    requirements_path = Path(directory) / "requirements.org"
    requirements = []
    if requirements_path.exists():
        requirements = parse_requirements(requirements_path)
        validate_requirements(requirements)
    requirements_by_id = {item["id"]: item for item in requirements}

    calibration_path = Path(directory) / "calibration.org"
    equipment_records = []
    if calibration_path.exists():
        equipment_records = parse_equipment(calibration_path)
        warnings.extend(validate_equipment(equipment_records))
    equipment_by_id = {item["id"]: item for item in equipment_records}
    reference_standard_ids = {
        standard for item in equipment_records for standard in item["standards"]
    }

    risk_path = Path(directory) / "risk.org"
    hazards = []
    if risk_path.exists():
        hazards = parse_hazards(risk_path)
        warnings.extend(validate_hazards(hazards))

    headings_by_document = []
    structured_items = {}
    design_items = {}
    equipment_items = {}
    acceptance_items = {}
    hazard_items = {}
    installation_items = {}
    service_items = {}
    user_items = {}
    for document in documents:
        for heading in read_headings(document["path"]):
            headings_by_document.append((document, heading))
            item_id = heading.properties.get("CUSTOM_ID", "").strip()
            if not item_id:
                continue
            if item_id in structured_items:
                raise ValueError(f"Duplicate structured item ID: {item_id}")
            item = {
                "id": item_id,
                "title": heading.title,
                "path": " / ".join([*heading.path, heading.title]),
                "document_id": document["metadata"]["ENGDOC_DOCUMENT_ID"],
                "requirements": split_references(
                    heading.properties.get(REFERENCE_PROPERTY, "")
                ),
                "design_items": split_references(
                    heading.properties.get("DESIGN_ITEMS", "")
                ),
                "hazards": split_references(heading.properties.get("HAZARDS", "")),
                "acceptance": split_references(
                    heading.properties.get("ACCEPTANCE_PROCEDURES", "")
                ),
                "equipment": split_references(heading.properties.get("EQUIPMENT", "")),
                "service_tasks": split_references(
                    heading.properties.get("SERVICE_TASKS", "")
                ),
            }
            structured_items[item_id] = item
            if item_id.startswith("DES-"):
                design_items[item_id] = item
            if item_id.startswith("EQ-"):
                equipment_items[item_id] = item
            if item_id.startswith("ATP-"):
                acceptance_items[item_id] = item
            if item_id.startswith("HAZ-"):
                hazard_items[item_id] = item
            if item_id.startswith("INS-"):
                installation_items[item_id] = item
            if item_id.startswith("SVC-"):
                service_items[item_id] = item
            if item_id.startswith("USR-"):
                user_items[item_id] = item

    links = {requirement_id: [] for requirement_id in requirements_by_id}
    design_links = {item_id: [] for item_id in design_items}
    equipment_links = {item_id: [] for item_id in equipment_items}
    for document, heading in headings_by_document:
        if document["path"] == requirements_path:
            continue
        references = split_references(heading.properties.get(REFERENCE_PROPERTY, ""))
        for requirement_id in references:
            if not requirements_path.exists():
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: requirement reference {requirement_id} "
                    "not checked because requirements.org is absent"
                )
                continue
            if requirement_id not in requirements_by_id:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown requirement {requirement_id}"
                )
            links[requirement_id].append({
                "document_id": document["metadata"]["ENGDOC_DOCUMENT_ID"],
                "item_id": heading.properties.get("CUSTOM_ID", ""),
                "heading": heading.title,
                "path": " / ".join([*heading.path, heading.title]),
            })

        for design_id in split_references(heading.properties.get("DESIGN_ITEMS", "")):
            if not design_items:
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"design reference {design_id} not checked because no design items exist"
                )
                continue
            if design_id not in design_items:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown design item {design_id}"
                )
            design_links[design_id].append({
                "document_id": document["metadata"]["ENGDOC_DOCUMENT_ID"],
                "item_id": heading.properties.get("CUSTOM_ID", ""),
                "path": " / ".join([*heading.path, heading.title]),
            })

        for equipment_id in split_references(heading.properties.get("EQUIPMENT", "")):
            if not equipment_items:
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"equipment reference {equipment_id} not checked because "
                    "calibration.org is absent"
                )
                continue
            if equipment_id not in equipment_items:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown equipment {equipment_id}"
                )
            record = equipment_by_id[equipment_id]
            if record["status"] != "ACTIVE":
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"equipment {equipment_id} is {record['status']}"
                )
            if record["due_date"] < date.today():
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"equipment {equipment_id} calibration is overdue"
                )
            equipment_links[equipment_id].append({
                "document_id": document["metadata"]["ENGDOC_DOCUMENT_ID"],
                "item_id": heading.properties.get("CUSTOM_ID", ""),
                "path": " / ".join([*heading.path, heading.title]),
            })

        for procedure_id in split_references(
            heading.properties.get("ACCEPTANCE_PROCEDURES", "")
        ):
            if not acceptance_items:
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"acceptance reference {procedure_id} not checked because "
                    "acceptance.org is absent"
                )
                continue
            if procedure_id not in acceptance_items:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown acceptance procedure {procedure_id}"
                )

        for hazard_id in split_references(heading.properties.get("HAZARDS", "")):
            if not hazard_items:
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"hazard reference {hazard_id} not checked because risk.org is absent"
                )
                continue
            if hazard_id not in hazard_items:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown hazard {hazard_id}"
                )

        for service_id in split_references(heading.properties.get("SERVICE_TASKS", "")):
            if not service_items:
                warnings.append(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / {heading.title}: "
                    f"service reference {service_id} not checked because service.org is absent"
                )
                continue
            if service_id not in service_items:
                raise ValueError(
                    f"{document['metadata']['ENGDOC_DOCUMENT_ID']} / "
                    f"{heading.title}: unknown service task {service_id}"
                )

    for requirement_id, targets in links.items():
        if not targets:
            warnings.append(f"{requirement_id}: no downstream references")
    for design_id, targets in design_links.items():
        if not targets:
            warnings.append(f"{design_id}: no acceptance references")
    for equipment_id, targets in equipment_links.items():
        if not targets and equipment_id not in reference_standard_ids:
            warnings.append(f"{equipment_id}: not assigned to an acceptance procedure")

    return {
        "documents": documents,
        "project_id": next(iter(project_ids)),
        "requirements": requirements,
        "links": links,
        "design_items": design_items,
        "design_links": design_links,
        "equipment": equipment_by_id,
        "equipment_links": equipment_links,
        "reference_standard_ids": reference_standard_ids,
        "hazards": hazards,
        "installation_items": installation_items,
        "service_items": service_items,
        "user_items": user_items,
        "warnings": warnings,
    }


def export_traceability(model, output):
    """Write MODEL to an XLSX traceability report at OUTPUT."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Traceability"
    headers = [
        "Requirement ID", "Requirement", "Category", "Verification",
        "Stage", "Referenced By", "Referenced Items", "Reference Paths",
        "Coverage",
    ]
    sheet.append(headers)
    for requirement in model["requirements"]:
        targets = model["links"][requirement["id"]]
        sheet.append([
            requirement["id"], requirement["requirement"], requirement["category"],
            requirement["verification"], requirement["stage"],
            ", ".join(sorted({target["document_id"] for target in targets})),
            ", ".join(target["item_id"] for target in targets if target["item_id"]),
            "\n".join(target["path"] for target in targets),
            "LINKED" if targets else "UNLINKED",
        ])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([22, 65, 32, 18, 14, 28, 30, 60, 14], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

    document_sheet = workbook.create_sheet("Documents")
    document_sheet.append(["Document ID", "Title", "Revision", "Status", "Related Documents"])
    for document in model["documents"]:
        metadata = document["metadata"]
        document_sheet.append([
            metadata["ENGDOC_DOCUMENT_ID"], metadata["TITLE"],
            metadata["ENGDOC_REVISION"], metadata["ENGDOC_STATUS"],
            ", ".join(document["related"]),
        ])
    for cell in document_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    document_sheet.freeze_panes = "A2"
    document_sheet.auto_filter.ref = document_sheet.dimensions
    for index, width in enumerate([28, 45, 12, 14, 55], 1):
        document_sheet.column_dimensions[get_column_letter(index)].width = width

    design_sheet = workbook.create_sheet("Design Coverage")
    design_sheet.append([
        "Design ID", "Title", "Document ID", "Requirements",
        "Acceptance Procedures", "Reference Paths", "Coverage",
    ])
    for design_id, item in model["design_items"].items():
        targets = model["design_links"][design_id]
        design_sheet.append([
            design_id, item["title"], item["document_id"],
            " ".join(item["requirements"]),
            ", ".join(target["item_id"] for target in targets if target["item_id"]),
            "\n".join(target["path"] for target in targets),
            "COVERED" if targets else "UNCOVERED",
        ])
    for cell in design_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in design_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([22, 36, 28, 34, 34, 60, 14], 1):
        design_sheet.column_dimensions[get_column_letter(index)].width = width
    design_sheet.freeze_panes = "B2"
    design_sheet.auto_filter.ref = design_sheet.dimensions

    equipment_sheet = workbook.create_sheet("Equipment Coverage")
    equipment_sheet.append([
        "Equipment ID", "Name", "Status", "Calibration Due", "Acceptance Procedures",
        "Reference Paths", "Usage",
    ])
    for equipment_id, item in model["equipment"].items():
        targets = model["equipment_links"][equipment_id]
        equipment_sheet.append([
            equipment_id, item["name"], item["status"], item["due_date"],
            ", ".join(target["item_id"] for target in targets if target["item_id"]),
            "\n".join(target["path"] for target in targets),
            "ASSIGNED" if targets else
            "REFERENCE STANDARD" if equipment_id in model["reference_standard_ids"]
            else "UNASSIGNED",
        ])
        equipment_sheet.cell(equipment_sheet.max_row, 4).number_format = "yyyy-mm-dd"
    for cell in equipment_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in equipment_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([20, 34, 18, 18, 34, 60, 14], 1):
        equipment_sheet.column_dimensions[get_column_letter(index)].width = width
    equipment_sheet.freeze_panes = "B2"
    equipment_sheet.auto_filter.ref = equipment_sheet.dimensions

    hazard_sheet = workbook.create_sheet("Hazard Coverage")
    hazard_sheet.append([
        "Hazard ID", "Title", "Status", "Initial Risk", "Residual Risk",
        "Requirements", "Design Items", "Acceptance Procedures", "Coverage",
    ])
    for hazard in model["hazards"]:
        covered = bool(
            hazard["requirements"] and hazard["design_items"] and hazard["acceptance"]
        )
        hazard_sheet.append([
            hazard["id"], hazard["title"], hazard["status"], hazard["initial_risk"],
            hazard["residual_risk"], " ".join(hazard["requirements"]),
            " ".join(hazard["design_items"]), " ".join(hazard["acceptance"]),
            "CONTROLLED" if covered else "INCOMPLETE",
        ])
    for cell in hazard_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in hazard_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([20, 36, 14, 14, 14, 32, 30, 34, 16], 1):
        hazard_sheet.column_dimensions[get_column_letter(index)].width = width
    hazard_sheet.freeze_panes = "B2"
    hazard_sheet.auto_filter.ref = hazard_sheet.dimensions

    installation_sheet = workbook.create_sheet("Installation Coverage")
    installation_sheet.append([
        "Installation ID", "Title", "Requirements", "Design Items", "Hazards",
        "Acceptance Procedures", "Equipment", "Coverage",
    ])
    for item_id, item in model["installation_items"].items():
        complete = all(
            item[field]
            for field in ("requirements", "design_items", "hazards", "acceptance")
        )
        installation_sheet.append([
            item_id, item["title"], " ".join(item["requirements"]),
            " ".join(item["design_items"]), " ".join(item["hazards"]),
            " ".join(item["acceptance"]), " ".join(item["equipment"]),
            "CONTROLLED" if complete else "INCOMPLETE",
        ])
    for cell in installation_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in installation_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([20, 36, 30, 28, 26, 32, 26, 16], 1):
        installation_sheet.column_dimensions[get_column_letter(index)].width = width
    installation_sheet.freeze_panes = "B2"
    installation_sheet.auto_filter.ref = installation_sheet.dimensions

    service_sheet = workbook.create_sheet("Service Coverage")
    service_sheet.append([
        "Service Task ID", "Title", "Requirements", "Design Items", "Hazards",
        "Equipment", "Coverage",
    ])
    for item_id, item in model["service_items"].items():
        complete = all(item[field] for field in ("requirements", "design_items", "hazards"))
        service_sheet.append([
            item_id, item["title"], " ".join(item["requirements"]),
            " ".join(item["design_items"]), " ".join(item["hazards"]),
            " ".join(item["equipment"]), "CONTROLLED" if complete else "INCOMPLETE",
        ])
    for cell in service_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in service_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([20, 36, 30, 28, 26, 26, 16], 1):
        service_sheet.column_dimensions[get_column_letter(index)].width = width
    service_sheet.freeze_panes = "B2"
    service_sheet.auto_filter.ref = service_sheet.dimensions

    user_sheet = workbook.create_sheet("User Guidance Coverage")
    user_sheet.append([
        "User Topic ID", "Title", "Requirements", "Design Items", "Hazards",
        "Service Tasks", "Coverage",
    ])
    for item_id, item in model["user_items"].items():
        complete = all(
            item[field]
            for field in ("requirements", "design_items", "hazards", "service_tasks")
        )
        user_sheet.append([
            item_id, item["title"], " ".join(item["requirements"]),
            " ".join(item["design_items"]), " ".join(item["hazards"]),
            " ".join(item["service_tasks"]),
            "CONTROLLED" if complete else "INCOMPLETE",
        ])
    for cell in user_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in user_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([20, 36, 30, 28, 26, 28, 16], 1):
        user_sheet.column_dimensions[get_column_letter(index)].width = width
    user_sheet.freeze_panes = "B2"
    user_sheet.auto_filter.ref = user_sheet.dimensions

    first = model["documents"][0]["metadata"]
    report_metadata = dict(first)
    report_metadata.update({
        "TITLE": "Requirements Traceability Matrix",
        "ENGDOC_DOCUMENT_ID": f"{model['project_id']}-TRACE",
        "ENGDOC_RELATED_DOCUMENTS": " ".join(
            item["metadata"]["ENGDOC_DOCUMENT_ID"] for item in model["documents"]
        ),
    })
    apply_workbook_metadata(workbook, report_metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("directory", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        model = validate_project_set(args.directory)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    for warning in model["warnings"]:
        print(f"WARNING: {warning}")
    if args.validate:
        print(
            f"Project-set validation passed: {len(model['documents'])} documents, "
            f"{len(model['warnings'])} warning(s)"
        )
    else:
        export_traceability(model, args.output)
        print(f"Exported traceability to {args.output}")


if __name__ == "__main__":
    main()
