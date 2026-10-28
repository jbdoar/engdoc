#!/usr/bin/env python3
"""Validate an engineering design Org document and export its register."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings


DESIGN_ID = re.compile(r"^DES-[A-Za-z0-9-]+$")
VALID_STATUSES = {"DRAFT", "REVIEW", "APPROVED", "SUPERSEDED"}


def parse_design_items(path):
    """Return structured design items from Org PATH."""
    items = []
    for heading in read_headings(path):
        design_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not design_id.startswith("DES-"):
            continue
        if not DESIGN_ID.fullmatch(design_id):
            raise ValueError(f"Invalid design ID: {design_id}")
        description = " ".join(
            line.strip()
            for line in heading.body
            if line.strip()
            and not line.lstrip().startswith("#")
            and not line.lstrip().startswith(":")
        )
        if not description:
            raise ValueError(f"{design_id}: missing design description")
        items.append({
            "id": design_id,
            "title": heading.title,
            "category": " / ".join(heading.path),
            "description": description,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "requirements": heading.properties.get("REQUIREMENTS", "").strip(),
        })

    seen = set()
    for item in items:
        if item["id"] in seen:
            raise ValueError(f"Duplicate design ID: {item['id']}")
        seen.add(item["id"])
    return items


def validate_design(items):
    """Validate ITEMS and return nonfatal warning messages."""
    if not items:
        raise ValueError("No design items found (expected DES-* CUSTOM_ID)")
    warnings = []
    for item in items:
        design_id = item["id"]
        if not item["status"]:
            raise ValueError(f"{design_id}: missing STATUS")
        if item["status"] not in VALID_STATUSES:
            raise ValueError(
                f"{design_id}: invalid STATUS {item['status']!r}; expected one of "
                f"{', '.join(sorted(VALID_STATUSES))}"
            )
        if not item["owner"]:
            raise ValueError(f"{design_id}: missing OWNER")
        if not item["requirements"]:
            warnings.append(f"{design_id}: no requirement references")
    return warnings


def export_xlsx(items, output, metadata):
    """Export design ITEMS and METADATA to an XLSX register."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Design Register"
    headers = [
        "Design ID", "Title", "Category", "Description", "Status", "Owner",
        "Requirements",
    ]
    sheet.append(headers)
    for item in items:
        sheet.append([
            item["id"], item["title"], item["category"], item["description"],
            item["status"], item["owner"], item["requirements"],
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([22, 34, 32, 72, 16, 28, 38], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        items = parse_design_items(args.input)
        warnings = validate_design(items)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(items)} design items, {len(warnings)} warning(s)")
    else:
        export_xlsx(items, args.output, metadata)
        print(f"Exported {len(items)} design items to {args.output}")


if __name__ == "__main__":
    main()
