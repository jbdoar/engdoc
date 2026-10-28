"""Shared document identity, validation, and workbook metadata helpers."""

from datetime import date
from pathlib import Path
import re

from openpyxl.styles import Alignment, Font, PatternFill

from engdoc_org import read_keywords


FIELDS = (
    ("TITLE", "Title"),
    ("ENGDOC_PROJECT_NAME", "Project Name"),
    ("ENGDOC_PROJECT_ID", "Project ID"),
    ("ENGDOC_DOCUMENT_ID", "Document ID"),
    ("ENGDOC_REVISION", "Revision"),
    ("ENGDOC_STATUS", "Status"),
    ("AUTHOR", "Author / Owner"),
    ("ENGDOC_ISSUE_DATE", "Issue Date"),
    ("ENGDOC_RELATED_DOCUMENTS", "Related Documents"),
)
REQUIRED = {
    "TITLE", "ENGDOC_PROJECT_NAME", "ENGDOC_PROJECT_ID",
    "ENGDOC_DOCUMENT_ID", "ENGDOC_REVISION", "ENGDOC_STATUS", "AUTHOR",
}
VALID_STATUSES = {"DRAFT", "REVIEW", "RELEASED", "OBSOLETE"}
IDENTIFIER = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")
REVISION = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")


def read_metadata(path):
    """Return normalized engdoc metadata from an Org document PATH."""
    keywords = read_keywords(path)
    return {key: keywords.get(key, "").strip() for key, _label in FIELDS}


def validate_metadata(metadata):
    """Validate METADATA and return nonfatal warning messages."""
    missing = [label for key, label in FIELDS if key in REQUIRED and not metadata.get(key)]
    if missing:
        raise ValueError(f"Missing document metadata: {', '.join(missing)}")

    for key, label in (
        ("ENGDOC_PROJECT_ID", "Project ID"),
        ("ENGDOC_DOCUMENT_ID", "Document ID"),
    ):
        if not IDENTIFIER.fullmatch(metadata[key]):
            raise ValueError(f"Invalid {label}: {metadata[key]!r}")

    if not REVISION.fullmatch(metadata["ENGDOC_REVISION"]):
        raise ValueError(f"Invalid Revision: {metadata['ENGDOC_REVISION']!r}")

    status = metadata["ENGDOC_STATUS"].upper()
    if status not in VALID_STATUSES:
        raise ValueError(
            f"Invalid Status {status!r}; expected one of "
            f"{', '.join(sorted(VALID_STATUSES))}"
        )
    metadata["ENGDOC_STATUS"] = status

    warnings = []
    issue_date = metadata.get("ENGDOC_ISSUE_DATE", "")
    if issue_date:
        try:
            date.fromisoformat(issue_date)
        except ValueError as exc:
            raise ValueError("Issue Date must use YYYY-MM-DD format") from exc
    elif status == "RELEASED":
        raise ValueError("Released documents require an Issue Date")
    else:
        warnings.append("draft/review document has no issue date")
    return warnings


def load_and_validate_metadata(path):
    """Read and validate metadata from PATH, returning metadata and warnings."""
    metadata = read_metadata(Path(path))
    warnings = validate_metadata(metadata)
    return metadata, warnings


def apply_workbook_metadata(workbook, metadata):
    """Apply METADATA to workbook properties and a visible Metadata sheet."""
    workbook.properties.title = metadata["TITLE"]
    workbook.properties.subject = metadata["ENGDOC_PROJECT_NAME"]
    workbook.properties.creator = metadata["AUTHOR"]
    workbook.properties.identifier = metadata["ENGDOC_DOCUMENT_ID"]
    workbook.properties.description = (
        f"Project {metadata['ENGDOC_PROJECT_ID']}; revision "
        f"{metadata['ENGDOC_REVISION']}; status {metadata['ENGDOC_STATUS']}"
    )

    if "Metadata" in workbook.sheetnames:
        del workbook["Metadata"]
    sheet = workbook.create_sheet("Metadata")
    sheet.append(["Field", "Value"])
    for key, label in FIELDS:
        sheet.append([label, metadata.get(key, "")])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 72
    sheet.freeze_panes = "A2"
