#!/usr/bin/env python3
"""Validate user-manual topics and export a controlled content register."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings

TOPIC_ID = re.compile(r"^USR-[A-Za-z0-9-]+$")
STATUSES = {"DRAFT", "REVIEW", "APPROVED", "RETIRED"}
TYPES = {"SAFETY", "STARTUP", "OPERATION", "SHUTDOWN", "RECOVERY", "CARE"}
AUDIENCES = {"OPERATOR", "SUPERVISOR", "OWNER"}


def split_values(value):
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def parse_topics(path):
    topics = []
    for heading in read_headings(path):
        topic_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not topic_id.startswith("USR-"):
            continue
        if not TOPIC_ID.fullmatch(topic_id):
            raise ValueError(f"Invalid user topic ID: {topic_id}")
        content = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not content:
            raise ValueError(f"{topic_id}: missing user guidance")
        topics.append({
            "id": topic_id, "title": heading.title,
            "category": " / ".join(heading.path), "content": content,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "type": heading.properties.get("TOPIC_TYPE", "").strip().upper(),
            "audience": heading.properties.get("AUDIENCE", "").strip().upper(),
            "requirements": split_values(heading.properties.get("REQUIREMENTS", "")),
            "design_items": split_values(heading.properties.get("DESIGN_ITEMS", "")),
            "hazards": split_values(heading.properties.get("HAZARDS", "")),
            "service_tasks": split_values(heading.properties.get("SERVICE_TASKS", "")),
        })
    seen = set()
    for topic in topics:
        if topic["id"] in seen:
            raise ValueError(f"Duplicate user topic ID: {topic['id']}")
        seen.add(topic["id"])
    return topics


def validate_topics(topics):
    if not topics:
        raise ValueError("No user topics found (expected USR-* CUSTOM_ID)")
    warnings = []
    for topic in topics:
        topic_id = topic["id"]
        if topic["status"] not in STATUSES:
            raise ValueError(f"{topic_id}: invalid or missing STATUS")
        if topic["type"] not in TYPES:
            raise ValueError(f"{topic_id}: invalid or missing TOPIC_TYPE")
        if topic["audience"] not in AUDIENCES:
            raise ValueError(f"{topic_id}: invalid or missing AUDIENCE")
        for field, label in (
            ("requirements", "requirements"), ("design_items", "design items"),
            ("hazards", "hazards"),
        ):
            if not topic[field]:
                warnings.append(f"{topic_id}: no {label} references")
        if topic["type"] in {"SAFETY", "RECOVERY"} and not topic["service_tasks"]:
            warnings.append(f"{topic_id}: no service escalation reference")
    return warnings


def export_xlsx(topics, output, metadata):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "User Content Register"
    sheet.append([
        "Topic ID", "Title", "Category", "Type", "Audience", "Status", "Guidance",
        "Requirements", "Design Items", "Hazards", "Service Tasks",
    ])
    for topic in topics:
        sheet.append([
            topic["id"], topic["title"], topic["category"], topic["type"],
            topic["audience"], topic["status"], topic["content"],
            " ".join(topic["requirements"]), " ".join(topic["design_items"]),
            " ".join(topic["hazards"]), " ".join(topic["service_tasks"]),
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([20, 34, 28, 16, 16, 14, 75, 30, 28, 26, 28], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "G2"
    sheet.auto_filter.ref = sheet.dimensions
    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        topics = parse_topics(args.input)
        warnings = validate_topics(topics)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(topics)} topics, {len(warnings)} warning(s)")
    else:
        export_xlsx(topics, args.output, metadata)
        print(f"Exported {len(topics)} user topics to {args.output}")


if __name__ == "__main__":
    main()
