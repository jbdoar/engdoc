#!/usr/bin/env python3
"""Export an Org requirements register to XLSX."""

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_org import read_headings
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata


REQUIREMENT_ID = re.compile(r"^REQ-[A-Za-z0-9-]+$")


def parse_requirements(path):
    """Return requirements extracted from an Org file."""
    requirements = []
    for heading in read_headings(path):
        props = heading.properties
        req_id = props.get("CUSTOM_ID", "")

        if not req_id.startswith("REQ-"):
            continue

        if not REQUIREMENT_ID.fullmatch(req_id):
            raise ValueError(f"Invalid requirement ID: {req_id}")

        statement = " ".join(
            line.strip()
            for line in heading.body
            if line.strip()
            and not line.lstrip().startswith("#")
            and not line.lstrip().startswith(":")
        )

        if not statement:
            raise ValueError(f"{req_id}: missing requirement statement")

        requirements.append({
            "id": req_id,
            "requirement": statement,
            "category": " / ".join(heading.path),
            "verification": props.get("VERIFY_METHOD", ""),
            "stage": props.get("VERIFY_STAGE", ""),
            "acceptance": props.get("ACCEPTANCE_CRITERIA", ""),
        })

    ids = [req["id"] for req in requirements]

    seen = set()

    for req_id in ids:
        if req_id in seen:
            raise ValueError(f"Duplicate requirement ID: {req_id}")
        seen.add(req_id)

    return requirements


VALID_METHODS = {
    "TEST",
    "INSPECTION",
    "ANALYSIS",
    "DEMONSTRATION",
}

VALID_STAGES = {
    "FAT",
    "SAT",
    "NONE",
}


def parse_stages(value):
    """Return the verification stages specified by a requirement."""
    stages = value.upper().replace(",", " ").split()

    unknown = set(stages) - VALID_STAGES

    if unknown:
        raise ValueError(
            "Unrecognized verification stage(s): "
            f"{', '.join(sorted(unknown))}"
        )

    if "NONE" in stages and len(stages) > 1:
        raise ValueError(
            "NONE cannot be combined with FAT or SAT"
        )

    return set(stages)


def validate_requirements(requirements):
    """Validate requirements and print nonfatal warnings."""
    if not requirements:
        raise ValueError(
            "No requirements found (expected REQ-* CUSTOM_ID)"
        )

    warnings = 0

    for req in requirements:
        req_id = req["id"]

        if not req["acceptance"]:
            print(
                f"WARNING: {req_id}: missing acceptance criteria",
                flush=True,
            )
            warnings += 1

        method = req["verification"].strip().upper()

        if not method:
            print(
                f"WARNING: {req_id}: missing verification method",
                flush=True,
            )
            warnings += 1

        elif method not in VALID_METHODS:
            print(
                f"WARNING: {req_id}: unrecognized verification "
                f"method: {req['verification']}",
                flush=True,
            )
            warnings += 1

        try:
            stages = parse_stages(req["stage"])
        except ValueError as exc:
            raise ValueError(f"{req_id}: {exc}") from exc

        if not stages:
            print(
                f"WARNING: {req_id}: missing verification stage",
                flush=True,
            )
            warnings += 1

    return warnings


def export_xlsx(requirements, output, metadata=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"

    headers = [
        "ID",
        "Requirement",
        "Category",
        "Verification",
        "Stage",
        "Acceptance Criteria",
    ]
    ws.append(headers)

    for req in requirements:
        ws.append([
            req["id"],
            req["requirement"],
            req["category"],
            req["verification"],
            req["stage"],
            req["acceptance"],
        ])

    # Header formatting.
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="243746",
        )
        cell.alignment = Alignment(wrap_text=True)

    widths = [22, 75, 38, 20, 16, 65]

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    if metadata:
        apply_workbook_metadata(wb, metadata)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--validate",
        action="store_true",
        help="Validate requirements without generating XLSX",
    )

    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")

    args = parser.parse_args()

    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")

    metadata, metadata_warnings = load_and_validate_metadata(args.input)
    requirements = parse_requirements(args.input)
    warnings = validate_requirements(requirements)
    for warning in metadata_warnings:
        print(f"WARNING: metadata: {warning}", flush=True)
    warnings += len(metadata_warnings)

    if args.validate:
        print(
            f"Validation passed: {len(requirements)} requirements, "
            f"{warnings} warning(s)"
        )
        return

    export_xlsx(requirements, args.output, metadata)

    print(
        f"Exported {len(requirements)} requirements "
        f"to {args.output}"
    )


if __name__ == "__main__":
    main()
