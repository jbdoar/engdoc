#!/usr/bin/env python3
"""Validate a service manual and export maintenance/record workbooks."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings

TASK_ID = re.compile(r"^SVC-[A-Za-z0-9-]+$")
STATUSES = {"DRAFT", "APPROVED", "RETIRED"}
TASK_TYPES = {"INSPECTION", "PREVENTIVE", "DIAGNOSTIC", "CORRECTIVE"}


def split_values(value):
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def parse_tasks(path):
    tasks = []
    for heading in read_headings(path):
        task_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not task_id.startswith("SVC-"):
            continue
        if not TASK_ID.fullmatch(task_id):
            raise ValueError(f"Invalid service task ID: {task_id}")
        procedure = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not procedure:
            raise ValueError(f"{task_id}: missing service procedure")
        tasks.append({
            "id": task_id, "title": heading.title,
            "category": " / ".join(heading.path), "procedure": procedure,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "type": heading.properties.get("TASK_TYPE", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "skill": heading.properties.get("SKILL_LEVEL", "").strip(),
            "interval": heading.properties.get("INTERVAL_DAYS", "").strip(),
            "criteria": heading.properties.get("ACCEPTANCE_CRITERIA", "").strip(),
            "requirements": split_values(heading.properties.get("REQUIREMENTS", "")),
            "design_items": split_values(heading.properties.get("DESIGN_ITEMS", "")),
            "hazards": split_values(heading.properties.get("HAZARDS", "")),
            "equipment": split_values(heading.properties.get("EQUIPMENT", "")),
        })
    seen = set()
    for task in tasks:
        if task["id"] in seen:
            raise ValueError(f"Duplicate service task ID: {task['id']}")
        seen.add(task["id"])
    return tasks


def validate_tasks(tasks):
    if not tasks:
        raise ValueError("No service tasks found (expected SVC-* CUSTOM_ID)")
    warnings = []
    for task in tasks:
        task_id = task["id"]
        if task["status"] not in STATUSES:
            raise ValueError(f"{task_id}: invalid or missing STATUS")
        if task["type"] not in TASK_TYPES:
            raise ValueError(f"{task_id}: invalid or missing TASK_TYPE")
        for field in ("owner", "skill", "criteria"):
            if not task[field]:
                raise ValueError(f"{task_id}: missing {field.upper()}")
        if task["interval"]:
            try:
                interval = int(task["interval"])
            except ValueError as exc:
                raise ValueError(f"{task_id}: INTERVAL_DAYS must be positive") from exc
            if interval <= 0:
                raise ValueError(f"{task_id}: INTERVAL_DAYS must be positive")
        elif task["type"] in {"INSPECTION", "PREVENTIVE"}:
            raise ValueError(f"{task_id}: scheduled task requires INTERVAL_DAYS")
        for field, label in (
            ("requirements", "requirements"), ("design_items", "design items"),
            ("hazards", "hazards"),
        ):
            if not task[field]:
                warnings.append(f"{task_id}: no {label} references")
    return warnings


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(wrap_text=True)


def export_xlsx(tasks, output, metadata):
    workbook = Workbook()
    readme = workbook.active
    readme.title = "READ ME"
    readme["A1"] = "BLANK SERVICE RECORD TEMPLATE — NOT A COMPLETED RECORD"
    readme["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    readme["A1"].fill = PatternFill("solid", fgColor="C00000")
    readme["A3"] = "Preserve completed work orders and evidence under records/service/."
    readme.column_dimensions["A"].width = 100

    register = workbook.create_sheet("Maintenance Register")
    register.append([
        "Task ID", "Title", "Type", "Status", "Owner", "Skill Level",
        "Interval (days)", "Acceptance Criteria", "Procedure", "Requirements",
        "Design Items", "Hazards", "Equipment",
    ])
    for task in tasks:
        register.append([
            task["id"], task["title"], task["type"], task["status"], task["owner"],
            task["skill"], int(task["interval"]) if task["interval"] else None,
            task["criteria"], task["procedure"], " ".join(task["requirements"]),
            " ".join(task["design_items"]), " ".join(task["hazards"]),
            " ".join(task["equipment"]),
        ])
    style_header(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([20, 34, 16, 14, 24, 24, 16, 55, 70, 30, 28, 26, 26], 1):
        register.column_dimensions[get_column_letter(index)].width = width
    register.freeze_panes = "H2"
    register.auto_filter.ref = register.dimensions

    schedule = workbook.create_sheet("Maintenance Schedule")
    schedule.append(["Task ID", "Title", "Interval (days)", "Last Completed",
                     "Next Due", "Assigned To", "Status"])
    for task in tasks:
        if task["interval"]:
            schedule.append([task["id"], task["title"], int(task["interval"]),
                             None, None, None, None])
    style_header(schedule)
    schedule.freeze_panes = "D2"

    record = workbook.create_sheet("Service Record")
    record.append([
        "Task ID", "Title", "Asset / Serial", "Technician", "Date", "Result",
        "As-Found Condition", "Work Performed", "Parts Used", "Measurements",
        "Evidence", "As-Left Condition", "Approval",
    ])
    for task in tasks:
        record.append([task["id"], task["title"]] + [None] * 11)
    style_header(record)
    result = DataValidation(type="list", formula1='"PASS,FAIL,DEFERRED,NOT RUN"')
    record.add_data_validation(result)
    result.add(f"F2:F{record.max_row}")
    for row in record.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        record.row_dimensions[row[0].row].height = 72
    for index, width in enumerate([20, 34, 24, 24, 16, 16, 40, 55, 34, 36, 30, 40, 24], 1):
        record.column_dimensions[get_column_letter(index)].width = width
    record.freeze_panes = "C2"

    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        tasks = parse_tasks(args.input)
        warnings = validate_tasks(tasks)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(tasks)} tasks, {len(warnings)} warning(s)")
    else:
        export_xlsx(tasks, args.output, metadata)
        print(f"Exported {len(tasks)} service tasks to {args.output}")


if __name__ == "__main__":
    main()
