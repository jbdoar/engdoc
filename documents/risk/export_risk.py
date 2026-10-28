#!/usr/bin/env python3
"""Validate a risk and hazard analysis and export its hazard register."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings


HAZARD_ID = re.compile(r"^HAZ-[A-Za-z0-9-]+$")
VALID_STATUSES = {"OPEN", "MITIGATED", "ACCEPTED", "CLOSED"}


def split_values(value):
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def score(value, label):
    try:
        number = int(value)
    except ValueError as exc:
        raise ValueError(f"{label} must be an integer from 1 to 5") from exc
    if not 1 <= number <= 5:
        raise ValueError(f"{label} must be an integer from 1 to 5")
    return number


def parse_hazards(path):
    """Return hazards parsed from Org PATH."""
    hazards = []
    for heading in read_headings(path):
        hazard_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not hazard_id.startswith("HAZ-"):
            continue
        if not HAZARD_ID.fullmatch(hazard_id):
            raise ValueError(f"Invalid hazard ID: {hazard_id}")
        description = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not description:
            raise ValueError(f"{hazard_id}: missing hazard and mitigation description")
        severity = score(heading.properties.get("SEVERITY", ""), f"{hazard_id} SEVERITY")
        likelihood = score(
            heading.properties.get("LIKELIHOOD", ""), f"{hazard_id} LIKELIHOOD"
        )
        residual_severity = score(
            heading.properties.get("RESIDUAL_SEVERITY", ""),
            f"{hazard_id} RESIDUAL_SEVERITY",
        )
        residual_likelihood = score(
            heading.properties.get("RESIDUAL_LIKELIHOOD", ""),
            f"{hazard_id} RESIDUAL_LIKELIHOOD",
        )
        hazards.append({
            "id": hazard_id,
            "title": heading.title,
            "category": " / ".join(heading.path),
            "description": description,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "severity": severity,
            "likelihood": likelihood,
            "initial_risk": severity * likelihood,
            "residual_severity": residual_severity,
            "residual_likelihood": residual_likelihood,
            "residual_risk": residual_severity * residual_likelihood,
            "requirements": split_values(heading.properties.get("REQUIREMENTS", "")),
            "design_items": split_values(heading.properties.get("DESIGN_ITEMS", "")),
            "acceptance": split_values(
                heading.properties.get("ACCEPTANCE_PROCEDURES", "")
            ),
        })
    seen = set()
    for hazard in hazards:
        if hazard["id"] in seen:
            raise ValueError(f"Duplicate hazard ID: {hazard['id']}")
        seen.add(hazard["id"])
    return hazards


def validate_hazards(hazards):
    """Validate HAZARDS and return nonfatal warnings."""
    if not hazards:
        raise ValueError("No hazards found (expected HAZ-* CUSTOM_ID)")
    warnings = []
    for hazard in hazards:
        hazard_id = hazard["id"]
        if hazard["status"] not in VALID_STATUSES:
            raise ValueError(
                f"{hazard_id}: invalid or missing STATUS; expected one of "
                f"{', '.join(sorted(VALID_STATUSES))}"
            )
        if not hazard["owner"]:
            raise ValueError(f"{hazard_id}: missing OWNER")
        if hazard["residual_risk"] > hazard["initial_risk"]:
            raise ValueError(f"{hazard_id}: residual risk exceeds initial risk")
        if not hazard["requirements"]:
            warnings.append(f"{hazard_id}: no requirement controls")
        if not hazard["design_items"]:
            warnings.append(f"{hazard_id}: no design controls")
        if not hazard["acceptance"]:
            warnings.append(f"{hazard_id}: no acceptance verification")
        if hazard["status"] in {"MITIGATED", "CLOSED"} and hazard["residual_risk"] >= 15:
            warnings.append(f"{hazard_id}: high residual risk is marked {hazard['status']}")
    return warnings


def risk_fill(value):
    if value >= 15:
        return "C00000"
    if value >= 8:
        return "FFC000"
    return "70AD47"


def export_xlsx(hazards, output, metadata):
    """Export HAZARDS to an XLSX hazard register and risk matrix."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hazard Register"
    headers = [
        "Hazard ID", "Title", "Category", "Status", "Owner", "Description / Controls",
        "Severity", "Likelihood", "Initial Risk", "Residual Severity",
        "Residual Likelihood", "Residual Risk", "Requirements", "Design Items",
        "Acceptance Procedures",
    ]
    sheet.append(headers)
    for hazard in hazards:
        sheet.append([
            hazard["id"], hazard["title"], hazard["category"], hazard["status"],
            hazard["owner"], hazard["description"], hazard["severity"],
            hazard["likelihood"], hazard["initial_risk"], hazard["residual_severity"],
            hazard["residual_likelihood"], hazard["residual_risk"],
            " ".join(hazard["requirements"]), " ".join(hazard["design_items"]),
            " ".join(hazard["acceptance"]),
        ])
        for column in (9, 12):
            cell = sheet.cell(sheet.max_row, column)
            cell.fill = PatternFill("solid", fgColor=risk_fill(cell.value))
            cell.font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(
        [20, 34, 28, 14, 24, 70, 10, 12, 12, 16, 18, 14, 32, 30, 34], 1
    ):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions

    matrix = workbook.create_sheet("Risk Matrix")
    matrix.append(["Likelihood \\ Severity", 1, 2, 3, 4, 5])
    for likelihood in range(5, 0, -1):
        matrix.append([likelihood] + [likelihood * severity for severity in range(1, 6)])
    for row in matrix.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=risk_fill(cell.value))
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
    for cell in matrix[1]:
        cell.font = Font(bold=True)
    matrix.column_dimensions["A"].width = 24
    for column in "BCDEF":
        matrix.column_dimensions[column].width = 12

    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        hazards = parse_hazards(args.input)
        warnings = validate_hazards(hazards)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(hazards)} hazards, {len(warnings)} warning(s)")
    else:
        export_xlsx(hazards, args.output, metadata)
        print(f"Exported {len(hazards)} hazards to {args.output}")


if __name__ == "__main__":
    main()
