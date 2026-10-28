#!/usr/bin/env python3
"""Validate calibration equipment and export a blank calibration workbook."""

import argparse
from datetime import date
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings


EQUIPMENT_ID = re.compile(r"^EQ-[A-Za-z0-9-]+$")
VALID_STATUSES = {"ACTIVE", "OUT-OF-SERVICE", "RETIRED"}


def split_values(value):
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def parse_equipment(path):
    """Return controlled equipment parsed from Org PATH."""
    equipment = []
    for heading in read_headings(path):
        equipment_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not equipment_id.startswith("EQ-"):
            continue
        if not EQUIPMENT_ID.fullmatch(equipment_id):
            raise ValueError(f"Invalid equipment ID: {equipment_id}")
        description = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not description:
            raise ValueError(f"{equipment_id}: missing calibration procedure")
        equipment.append({
            "id": equipment_id,
            "name": heading.title,
            "category": " / ".join(heading.path),
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "manufacturer": heading.properties.get("MANUFACTURER", "").strip(),
            "model": heading.properties.get("MODEL", "").strip(),
            "serial": heading.properties.get("SERIAL", "").strip(),
            "range": heading.properties.get("RANGE", "").strip(),
            "accuracy": heading.properties.get("ACCURACY", "").strip(),
            "interval": heading.properties.get("INTERVAL_DAYS", "").strip(),
            "due": heading.properties.get("CALIBRATION_DUE", "").strip(),
            "standards": split_values(
                heading.properties.get("REFERENCE_STANDARDS", "")
            ),
            "procedure": description,
        })
    seen = set()
    for item in equipment:
        if item["id"] in seen:
            raise ValueError(f"Duplicate equipment ID: {item['id']}")
        seen.add(item["id"])
    return equipment


def validate_equipment(equipment, today=None):
    """Validate EQUIPMENT and return nonfatal validity warnings."""
    if not equipment:
        raise ValueError("No calibration equipment found (expected EQ-* CUSTOM_ID)")
    today = today or date.today()
    by_id = {item["id"]: item for item in equipment}
    warnings = []
    for item in equipment:
        equipment_id = item["id"]
        if item["status"] not in VALID_STATUSES:
            raise ValueError(
                f"{equipment_id}: invalid or missing STATUS; expected one of "
                f"{', '.join(sorted(VALID_STATUSES))}"
            )
        for field in ("owner", "manufacturer", "model", "serial", "range", "accuracy"):
            if not item[field]:
                raise ValueError(f"{equipment_id}: missing {field.upper()}")
        try:
            interval = int(item["interval"])
        except ValueError as exc:
            raise ValueError(f"{equipment_id}: INTERVAL_DAYS must be a positive integer") from exc
        if interval <= 0:
            raise ValueError(f"{equipment_id}: INTERVAL_DAYS must be a positive integer")
        try:
            due = date.fromisoformat(item["due"])
        except ValueError as exc:
            raise ValueError(f"{equipment_id}: CALIBRATION_DUE must use YYYY-MM-DD") from exc
        item["due_date"] = due
        if item["status"] == "ACTIVE" and due < today:
            warnings.append(f"{equipment_id}: calibration overdue since {due}")
        for standard in item["standards"]:
            if standard == equipment_id:
                raise ValueError(f"{equipment_id}: cannot reference itself as a standard")
            if standard not in by_id:
                raise ValueError(f"{equipment_id}: unknown reference standard {standard}")

    def visit(equipment_id, path):
        if equipment_id in path:
            raise ValueError(
                "Reference-standard cycle: " + " -> ".join([*path, equipment_id])
            )
        for standard in by_id[equipment_id]["standards"]:
            visit(standard, [*path, equipment_id])

    for equipment_id in by_id:
        visit(equipment_id, [])
    return warnings


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_xlsx(equipment, output, metadata):
    """Export equipment register and blank calibration execution template."""
    workbook = Workbook()
    readme = workbook.active
    readme.title = "READ ME"
    readme["A1"] = "BLANK CALIBRATION TEMPLATE — NOT A COMPLETED RECORD"
    readme["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    readme["A1"].fill = PatternFill("solid", fgColor="C00000")
    readme["A3"] = (
        "Save completed calibrations and certificates under records/calibration/. "
        "This generated workbook may be replaced during export."
    )
    readme["A3"].alignment = Alignment(wrap_text=True)
    readme.column_dimensions["A"].width = 100

    register = workbook.create_sheet("Equipment Register")
    headers = [
        "Equipment ID", "Name", "Category", "Status", "Owner", "Manufacturer",
        "Model", "Serial", "Range", "Accuracy", "Interval (days)",
        "Calibration Due", "Reference Standards", "Procedure",
    ]
    register.append(headers)
    for item in equipment:
        register.append([
            item["id"], item["name"], item["category"], item["status"],
            item["owner"], item["manufacturer"], item["model"], item["serial"],
            item["range"], item["accuracy"], int(item["interval"]),
            item["due_date"], " ".join(item["standards"]), item["procedure"],
        ])
        register.cell(register.max_row, 12).number_format = "yyyy-mm-dd"
    style_header(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(
        [20, 32, 28, 18, 24, 26, 20, 20, 28, 30, 16, 18, 28, 65], 1
    ):
        register.column_dimensions[get_column_letter(index)].width = width
    register.freeze_panes = "D2"
    register.auto_filter.ref = register.dimensions

    execution = workbook.create_sheet("Calibration Execution")
    execution_headers = [
        "Equipment ID", "Name", "Serial", "Calibration Date", "Technician",
        "Reference Standards / Certificate IDs", "As-Found Result",
        "Adjustment / Repair", "As-Left Result", "Next Due Date", "Result",
        "Certificate / Evidence Reference", "Approval",
    ]
    execution.append(execution_headers)
    for item in equipment:
        execution.append([item["id"], item["name"], item["serial"]] + [None] * 10)
    style_header(execution)
    for row in execution.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        execution.row_dimensions[row[0].row].height = 72
    result_validation = DataValidation(
        type="list", formula1='"PASS,FAIL,LIMITED,NOT RUN"', allow_blank=True
    )
    execution.add_data_validation(result_validation)
    result_validation.add(f"K2:K{execution.max_row}")
    for index, width in enumerate([20, 30, 20, 18, 24, 40, 36, 36, 36, 18, 16, 38, 24], 1):
        execution.column_dimensions[get_column_letter(index)].width = width
    execution.freeze_panes = "D2"
    execution.auto_filter.ref = execution.dimensions

    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        equipment = parse_equipment(args.input)
        warnings = validate_equipment(equipment)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(equipment)} items, {len(warnings)} warning(s)")
    else:
        export_xlsx(equipment, args.output, metadata)
        print(f"Exported {len(equipment)} equipment items to {args.output}")


if __name__ == "__main__":
    main()
