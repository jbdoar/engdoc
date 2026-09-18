#!/usr/bin/env python3

# ~/.emacs.d/python/org/export_schedule.py

import argparse
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


TODO_STATES = (
    r"TODO|SPECIFY|QUOTED|ORDERED|RECEIVED|"
    r"NEXT|WAITING|DONE|CANCELLED"
)

HEADING_RE = re.compile(
    rf"^(?P<stars>\*+)\s+"
    rf"(?:(?P<todo>{TODO_STATES})\s+)?"
    rf"(?P<title>.*?)"
    rf"(?:\s+:(?P<tags>[^:]+(?::[^:]+)*):)?\s*$"
)

DATE_RE = re.compile(
    r"<(\d{4}-\d{2}-\d{2})(?:\s+[^>]*)?>"
)


def parse_date(text):
    match = DATE_RE.search(text)

    if not match:
        return None

    return datetime.strptime(
        match.group(1),
        "%Y-%m-%d",
    ).date()


def parse_project(path):
    lines = Path(path).read_text(
        encoding="utf-8"
    ).splitlines()

    entries = []

    for i, line in enumerate(lines):
        match = HEADING_RE.match(line)

        if not match:
            continue

        level = len(match.group("stars"))
        title = match.group("title").strip()
        todo = match.group("todo") or ""

        tags = [
            tag
            for tag in (match.group("tags") or "").split(":")
            if tag
        ]

        owners = [
            tag[1:]
            for tag in tags
            if tag.startswith("@") and len(tag) > 1
        ]

        scheduled = None
        deadline = None

        j = i + 1

        while (
            j < len(lines)
            and not HEADING_RE.match(lines[j])
        ):
            text = lines[j].strip()

            if "SCHEDULED:" in text:
                scheduled = parse_date(
                    text.split("SCHEDULED:", 1)[1]
                )

            if "DEADLINE:" in text:
                deadline = parse_date(
                    text.split("DEADLINE:", 1)[1]
                )

            j += 1

        entries.append({
            "LEVEL": level,
            "TASK": title,
            "TODO": todo,
            "TAGS": tags,
            "OWNERS": owners,
            "START": scheduled,
            "END": deadline,
        })

    keep = set()

    for i, entry in enumerate(entries):
        active_todo = (
            entry["TODO"]
            and entry["TODO"]
            not in {"DONE", "CANCELLED"}
        )

        relevant = (
            active_todo
            or entry["START"]
            or entry["END"]
        )

        if not relevant:
            continue

        keep.add(i)

        level = entry["LEVEL"]

        for parent_i in range(i - 1, -1, -1):
            parent = entries[parent_i]

            if parent["LEVEL"] < level:
                keep.add(parent_i)
                level = parent["LEVEL"]

                if level == 1:
                    break

    return [
        entry
        for i, entry in enumerate(entries)
        if i in keep
    ]


def write_gantt(entries, path):
    dated = [
        entry
        for entry in entries
        if entry["START"] or entry["END"]
    ]

    # Normalize entries having only one date.
    for entry in dated:
        entry["START"] = (
            entry["START"] or entry["END"]
        )

        entry["END"] = (
            entry["END"] or entry["START"]
        )

    # One timeline column per calendar day.
    days = []

    if dated:
        first_day = min(
            entry["START"]
            for entry in dated
        )

        last_day = max(
            entry["END"]
            for entry in dated
        )

        day = first_day

        while day <= last_day:
            days.append(day)
            day += timedelta(days=1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    # Preserve the Org hierarchy using Excel outline levels.
    ws.sheet_properties.outlinePr.summaryBelow = False

    headers = [
        "Task",
        "Owner",
        "Status",
        "Start",
        "End",
    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):
        cell = ws.cell(
            row=1,
            column=col,
            value=header,
        )

        cell.font = Font(bold=True)

    timeline_start_col = len(headers) + 1

    # Daily timeline headers.
    for col, day in enumerate(
        days,
        start=timeline_start_col,
    ):
        cell = ws.cell(
            row=1,
            column=col,
            value=day,
        )

        cell.number_format = "ddd m/d"
        cell.alignment = Alignment(
            text_rotation=90
        )
        cell.font = Font(bold=True)

    # Styles
    task_fill = PatternFill(
        "solid",
        fgColor="5B9BD5",
    )

    milestone_fill = PatternFill(
        "solid",
        fgColor="FFC000",
    )

    done_fill = PatternFill(
        "solid",
        fgColor="A9D18E",
    )

    overdue_fill = PatternFill(
        "solid",
        fgColor="F4CCCC",
    )

    group_fill = PatternFill(
        "solid",
        fgColor="D9E1F2",
    )

    today = date.today()

    today_border = Border(
        left=Side(
            style="medium",
            color="C00000",
        ),
        right=Side(
            style="medium",
            color="C00000",
        ),
    )

    row_num = 2

    for entry in entries:
        level = entry["LEVEL"]
        start = entry["START"]
        end = entry["END"]
        todo = entry["TODO"]

        task_cell = ws.cell(
            row=row_num,
            column=1,
            value=entry["TASK"],
        )

        task_cell.alignment = Alignment(
            indent=max(level - 1, 0)
        )

        ws.row_dimensions[
            row_num
        ].outlineLevel = min(
            max(level - 1, 0),
            7,
        )

        ws.cell(
            row=row_num,
            column=2,
            value=", ".join(entry["OWNERS"]),
        )

        # Context/group heading.
        if not todo and not start and not end:
            task_cell.font = Font(bold=True)

            for col in range(
                1,
                timeline_start_col + len(days),
            ):
                ws.cell(
                    row=row_num,
                    column=col,
                ).fill = group_fill

            row_num += 1
            continue

        ws.cell(
            row=row_num,
            column=3,
            value=todo,
        )

        # Active TODO with no dates yet.
        if not start and not end:
            ws.cell(
                row=row_num,
                column=3,
                value="UNSCHEDULED",
            )

            row_num += 1
            continue

        start = start or end
        end = end or start

        ws.cell(
            row=row_num,
            column=4,
            value=start,
        ).number_format = "mmm d, yyyy"

        ws.cell(
            row=row_num,
            column=5,
            value=end,
        ).number_format = "mmm d, yyyy"

        milestone = start == end
        done = todo == "DONE"
        overdue = end < today and not done

        for day_index, day in enumerate(days):
            cell = ws.cell(
                row=row_num,
                column=timeline_start_col + day_index,
            )

            # Mark today's column.
            if day == today:
                cell.border = today_border

            if start <= day <= end:
                cell.fill = (
                    done_fill
                    if done
                    else overdue_fill
                    if overdue
                    else milestone_fill
                    if milestone
                    else task_fill
                )

        row_num += 1

    ws.freeze_panes = "F2"

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14

    # Narrow daily timeline columns.
    for col in range(
        timeline_start_col,
        timeline_start_col + len(days),
    ):
        ws.column_dimensions[
            ws.cell(
                row=1,
                column=col,
            ).column_letter
        ].width = 4

    try:
        wb.save(path)

    except PermissionError:
        raise SystemExit(
            f"Cannot write {path}: the file may be open in Excel. "
            "Close it and run the export again."
        )


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "input",
        help="Org project file",
    )

    parser.add_argument(
        "output",
        help="Output schedule.xlsx file",
    )

    args = parser.parse_args()

    entries = parse_project(args.input)
    write_gantt(
        entries,
        Path(args.output),
    )


if __name__ == "__main__":
    main()
