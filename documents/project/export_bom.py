#!/usr/bin/env python3

# ~/.emacs.d/python/org/export_bom.py

import argparse
import csv
import importlib.util
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_org import read_headings
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata


TODO_STATES = r"TODO|SPECIFY|QUOTED|ORDERED|RECEIVED|WAITING|DONE|CANCELLED"

ORG_LINK_RE = re.compile(
    r"^\[\[(?P<target>[^\]]+)\](?:\[(?P<label>[^\]]*)\])?\]$"
)


FIELDS = [
    "LEVEL",
    "ITEM",
    "STATE",
    "QTY",
    "PN",
    "MPN",
    "MFR",
    "VENDOR",
    "VENDOR_PN",
    "SOURCE",
    "COST",
    "URL",
]


def parse_cost(value):
    return float(
        str(value)
        .replace("$", "")
        .replace(",", "")
        .strip()
    )


NUMERIC_FIELDS = {
    "LEVEL": int,
    "QTY": int,
    "COST": parse_cost,
}


def parse_bom(path):
    rows = []
    states = TODO_STATES.split("|")

    for heading in read_headings(path, states):
        if "BOM" not in heading.tags:
            continue

        row = {
            "LEVEL": heading.level,
            "ITEM": heading.title,
            "STATE": heading.todo,
        }

        for field in FIELDS[3:]:
            row[field] = heading.properties.get(field, "")

        rows.append(row)

    return rows


def parse_org_link(value):
    value = value.strip()
    match = ORG_LINK_RE.match(value)

    if not match:
        return value, None

    target = match.group("target")
    label = match.group("label") or target

    return label, target


def write_csv(rows, path):
    try:
        with open(
            path,
            "w",
            newline="",
            encoding="utf-8-sig",
        ) as f:
            writer = csv.DictWriter(
                f,
                fieldnames=FIELDS,
            )

            writer.writeheader()

            for row in rows:
                output_row = {}

                for field in FIELDS:
                    value = str(row.get(field, ""))
                    label, target = parse_org_link(value)

                    output_row[field] = (
                        target if target else value
                    )

                writer.writerow(output_row)

    except PermissionError:
        raise SystemExit(
            f"Cannot write {path}: the file may be open "
            "in another application. Close it and run "
            "the export again."
        )


def write_xlsx(rows, path, metadata=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"

    # Preserve Org hierarchy using Excel row outlines.
    ws.sheet_properties.outlinePr.summaryBelow = False

    # Header
    for column_number, field in enumerate(
        FIELDS,
        start=1,
    ):
        cell = ws.cell(
            row=1,
            column=column_number,
            value=field,
        )
        cell.font = Font(bold=True)

    # Data
    for row_number, row in enumerate(
        rows,
        start=2,
    ):
        level = int(row["LEVEL"])

        ws.row_dimensions[row_number].outlineLevel = min(
            max(level - 1, 0),
            7,
        )

        for column_number, field in enumerate(
            FIELDS,
            start=1,
        ):
            value = row.get(field, "")
            target = None

            # Convert only fields that should genuinely be numeric.
            if field in NUMERIC_FIELDS and value not in ("", None):
                try:
                    value = NUMERIC_FIELDS[field](value)
                except (ValueError, TypeError):
                    pass

            # Everything else stays textual, including PN/MPN/vendor PNs.
            if isinstance(value, str):
                value, target = parse_org_link(value)

            cell = ws.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

            if field == "QTY" and isinstance(value, int):
                cell.number_format = "0"

            if field == "COST" and isinstance(value, (int, float)):
                cell.number_format = "$#,##0.00"

            if target:
                cell.hyperlink = target
                cell.style = "Hyperlink"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if metadata:
        apply_workbook_metadata(wb, metadata)

    # Auto-size columns.
    for column in ws.columns:
        width = max(
            len(str(cell.value or ""))
            for cell in column
        )

        ws.column_dimensions[
            column[0].column_letter
        ].width = min(width + 2, 40)

    try:
        wb.save(path)

    except PermissionError:
        raise SystemExit(
            f"Cannot write {path}: the file may be open "
            "in Excel. Close it and run the export again."
        )


def load_adapter(name, directories):
    """Load a named external BOM adapter from configured DIRECTORIES."""
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", name):
        raise ValueError(f"Invalid BOM format name: {name}")
    for directory in directories:
        candidate = Path(directory) / "bom" / name / "adapter.py"
        if not candidate.is_file():
            continue
        spec = importlib.util.spec_from_file_location(
            f"engdoc_bom_adapter_{name.replace('-', '_')}", candidate
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        export = getattr(module, "export", None)
        if not callable(export):
            raise ValueError(f"BOM adapter has no callable export(): {candidate}")
        return export
    raise ValueError(f"Unknown BOM format adapter: {name}")


def export_bom(rows, output, metadata, format_name="default", format_directories=()):
    """Export normalized ROWS using the selected external format contract."""
    output = Path(output)
    if format_name == "default":
        if output.suffix.lower() == ".csv":
            write_csv(rows, output)
        elif output.suffix.lower() == ".xlsx":
            write_xlsx(rows, output, metadata)
        else:
            raise ValueError("output must end in .csv or .xlsx")
        return
    adapter = load_adapter(format_name, format_directories)
    adapter(rows, output, metadata)
    if not output.is_file() or output.stat().st_size == 0:
        raise ValueError(f"BOM adapter did not produce output: {output}")


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument("--format", default="default", dest="format_name")
    parser.add_argument(
        "--format-directory", action="append", default=[], dest="format_directories"
    )

    parser.add_argument(
        "input",
        help="Org project/BOM file",
    )

    parser.add_argument(
        "output",
        help="Output .xlsx or .csv file",
    )

    args = parser.parse_args()

    rows = parse_bom(args.input)
    metadata, _warnings = load_and_validate_metadata(args.input)
    try:
        export_bom(
            rows, Path(args.output), metadata,
            args.format_name, args.format_directories,
        )
    except ValueError as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    main()
