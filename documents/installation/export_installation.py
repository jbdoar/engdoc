#!/usr/bin/env python3
"""Validate installation activities and export a blank execution workbook."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings

ITEM_ID = re.compile(r"^INS-[A-Za-z0-9-]+$")
STATUSES = {"DRAFT", "READY", "COMPLETE", "BLOCKED"}
PHASES = {"SITE", "RECEIVING", "INSTALLATION", "COMMISSIONING", "HANDOVER"}


def split_values(value):
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def parse_activities(path):
    activities = []
    for heading in read_headings(path):
        item_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not item_id.startswith("INS-"):
            continue
        if not ITEM_ID.fullmatch(item_id):
            raise ValueError(f"Invalid installation ID: {item_id}")
        description = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not description:
            raise ValueError(f"{item_id}: missing activity description")
        activities.append({
            "id": item_id, "title": heading.title,
            "category": " / ".join(heading.path), "description": description,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "phase": heading.properties.get("PHASE", "").strip().upper(),
            "criteria": heading.properties.get("ACCEPTANCE_CRITERIA", "").strip(),
            "requirements": split_values(heading.properties.get("REQUIREMENTS", "")),
            "design_items": split_values(heading.properties.get("DESIGN_ITEMS", "")),
            "hazards": split_values(heading.properties.get("HAZARDS", "")),
            "acceptance": split_values(
                heading.properties.get("ACCEPTANCE_PROCEDURES", "")
            ),
            "equipment": split_values(heading.properties.get("EQUIPMENT", "")),
        })
    seen = set()
    for item in activities:
        if item["id"] in seen:
            raise ValueError(f"Duplicate installation ID: {item['id']}")
        seen.add(item["id"])
    return activities


def validate_activities(activities):
    if not activities:
        raise ValueError("No installation activities found (expected INS-* CUSTOM_ID)")
    warnings = []
    for item in activities:
        item_id = item["id"]
        if item["status"] not in STATUSES:
            raise ValueError(f"{item_id}: invalid or missing STATUS")
        if item["phase"] not in PHASES:
            raise ValueError(f"{item_id}: invalid or missing PHASE")
        if not item["owner"]:
            raise ValueError(f"{item_id}: missing OWNER")
        if not item["criteria"]:
            raise ValueError(f"{item_id}: missing ACCEPTANCE_CRITERIA")
        for field, label in (
            ("requirements", "requirements"), ("design_items", "design items"),
            ("hazards", "hazards"), ("acceptance", "acceptance procedures"),
        ):
            if not item[field]:
                warnings.append(f"{item_id}: no {label} references")
    return warnings


def header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(wrap_text=True)


def export_xlsx(activities, output, metadata):
    workbook = Workbook()
    readme = workbook.active
    readme.title = "READ ME"
    readme["A1"] = "BLANK SITE EXECUTION TEMPLATE — NOT A COMPLETED RECORD"
    readme["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    readme["A1"].fill = PatternFill("solid", fgColor="C00000")
    readme["A3"] = "Save completed site records beneath records/installation/."
    readme.column_dimensions["A"].width = 100

    register = workbook.create_sheet("Activity Register")
    fields = [
        "ID", "Title", "Phase", "Status", "Owner", "Criteria", "Description",
        "Requirements", "Design Items", "Hazards", "Acceptance Procedures", "Equipment",
    ]
    register.append(fields)
    for item in activities:
        register.append([
            item["id"], item["title"], item["phase"], item["status"], item["owner"],
            item["criteria"], item["description"], " ".join(item["requirements"]),
            " ".join(item["design_items"]), " ".join(item["hazards"]),
            " ".join(item["acceptance"]), " ".join(item["equipment"]),
        ])
    header(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([20, 34, 18, 14, 24, 60, 65, 30, 28, 26, 32, 26], 1):
        register.column_dimensions[get_column_letter(index)].width = width
    register.freeze_panes = "F2"
    register.auto_filter.ref = register.dimensions

    execution = workbook.create_sheet("Site Execution")
    execution.append([
        "ID", "Title", "Phase", "Acceptance Criteria", "Executor", "Date",
        "Configuration / Location", "Result", "Observation / Measurement",
        "Evidence", "Punch Item", "Disposition", "Approval",
    ])
    for item in activities:
        execution.append([item["id"], item["title"], item["phase"], item["criteria"]]
                         + [None] * 9)
    header(execution)
    result = DataValidation(type="list", formula1='"PASS,FAIL,BLOCKED,NOT RUN"')
    execution.add_data_validation(result)
    result.add(f"H2:H{execution.max_row}")
    for row in execution.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        execution.row_dimensions[row[0].row].height = 72
    for index, width in enumerate([20, 34, 18, 55, 24, 16, 32, 16, 48, 32, 24, 40, 24], 1):
        execution.column_dimensions[get_column_letter(index)].width = width
    execution.freeze_panes = "E2"
    execution.auto_filter.ref = execution.dimensions

    punch = workbook.create_sheet("Punch List")
    punch.append(["Punch ID", "Related Activity", "Description", "Owner", "Priority",
                  "Target Date", "Status", "Disposition", "Evidence", "Approval"])
    header(punch)
    for index in range(1, 21):
        punch.append([f"PUNCH-{index:03d}"] + [None] * 9)
    status = DataValidation(type="list", formula1='"OPEN,IN WORK,READY TO VERIFY,CLOSED"')
    punch.add_data_validation(status)
    status.add("G2:G21")
    for index, width in enumerate([18, 22, 55, 24, 14, 16, 20, 45, 30, 24], 1):
        punch.column_dimensions[get_column_letter(index)].width = width
    punch.freeze_panes = "C2"

    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        activities = parse_activities(args.input)
        warnings = validate_activities(activities)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(f"Validation passed: {len(activities)} activities, {len(warnings)} warning(s)")
    else:
        export_xlsx(activities, args.output, metadata)
        print(f"Exported {len(activities)} activities to {args.output}")


if __name__ == "__main__":
    main()
