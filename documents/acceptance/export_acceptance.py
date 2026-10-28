#!/usr/bin/env python3
"""Validate acceptance procedures and export a blank execution workbook."""

import argparse
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engdoc_metadata import apply_workbook_metadata, load_and_validate_metadata
from engdoc_org import read_headings


PROCEDURE_ID = re.compile(r"^ATP-[A-Za-z0-9-]+$")
VALID_STATUSES = {"DRAFT", "REVIEW", "APPROVED", "RETIRED"}
VALID_STAGES = {"FAT", "SAT", "COMMISSIONING"}


def split_values(value):
    """Split a comma- or whitespace-separated property value."""
    return list(dict.fromkeys(item for item in re.split(r"[\s,]+", value.strip()) if item))


def parse_procedures(path):
    """Return acceptance procedures parsed from Org PATH."""
    procedures = []
    for heading in read_headings(path):
        procedure_id = heading.properties.get("CUSTOM_ID", "").strip()
        if not procedure_id.startswith("ATP-"):
            continue
        if not PROCEDURE_ID.fullmatch(procedure_id):
            raise ValueError(f"Invalid acceptance procedure ID: {procedure_id}")
        description = " ".join(
            line.strip() for line in heading.body
            if line.strip() and not line.lstrip().startswith(("#", ":"))
        )
        if not description:
            raise ValueError(f"{procedure_id}: missing procedure description")
        procedures.append({
            "id": procedure_id,
            "title": heading.title,
            "category": " / ".join(heading.path),
            "description": description,
            "status": heading.properties.get("STATUS", "").strip().upper(),
            "owner": heading.properties.get("OWNER", "").strip(),
            "stages": split_values(heading.properties.get("STAGES", "").upper()),
            "requirements": split_values(heading.properties.get("REQUIREMENTS", "")),
            "design_items": split_values(heading.properties.get("DESIGN_ITEMS", "")),
            "equipment": split_values(heading.properties.get("EQUIPMENT", "")),
            "criteria": heading.properties.get("ACCEPTANCE_CRITERIA", "").strip(),
        })
    seen = set()
    for procedure in procedures:
        if procedure["id"] in seen:
            raise ValueError(f"Duplicate acceptance procedure ID: {procedure['id']}")
        seen.add(procedure["id"])
    return procedures


def validate_procedures(procedures):
    """Validate PROCEDURES and return nonfatal warnings."""
    if not procedures:
        raise ValueError("No acceptance procedures found (expected ATP-* CUSTOM_ID)")
    warnings = []
    for procedure in procedures:
        procedure_id = procedure["id"]
        if procedure["status"] not in VALID_STATUSES:
            raise ValueError(
                f"{procedure_id}: invalid or missing STATUS; expected one of "
                f"{', '.join(sorted(VALID_STATUSES))}"
            )
        if not procedure["owner"]:
            raise ValueError(f"{procedure_id}: missing OWNER")
        if not procedure["stages"]:
            raise ValueError(f"{procedure_id}: missing STAGES")
        unknown = set(procedure["stages"]) - VALID_STAGES
        if unknown:
            raise ValueError(
                f"{procedure_id}: invalid stage(s): {', '.join(sorted(unknown))}"
            )
        if not procedure["criteria"]:
            raise ValueError(f"{procedure_id}: missing ACCEPTANCE_CRITERIA")
        if not procedure["requirements"]:
            warnings.append(f"{procedure_id}: no requirement references")
        if not procedure["design_items"]:
            warnings.append(f"{procedure_id}: no design-item references")
        if not procedure["equipment"]:
            warnings.append(f"{procedure_id}: no equipment references")
    return warnings


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_xlsx(procedures, output, metadata):
    """Export PROCEDURES as a blank, explicitly disposable execution template."""
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "READ ME"
    instructions["A1"] = "BLANK EXECUTION TEMPLATE — NOT A COMPLETED RECORD"
    instructions["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    instructions["A1"].fill = PatternFill("solid", fgColor="C00000")
    instructions["A3"] = (
        "Before execution, save a uniquely named copy beneath records/. "
        "Exports are regenerated and may overwrite this blank template."
    )
    instructions["A5"] = "Source document"
    instructions["B5"] = metadata["ENGDOC_DOCUMENT_ID"]
    instructions["A6"] = "Revision"
    instructions["B6"] = metadata["ENGDOC_REVISION"]
    instructions.column_dimensions["A"].width = 34
    instructions.column_dimensions["B"].width = 90
    instructions["A3"].alignment = Alignment(wrap_text=True)

    register = workbook.create_sheet("Procedures")
    register_headers = [
        "Procedure ID", "Title", "Category", "Status", "Owner", "Stages",
        "Requirements", "Design Items", "Equipment", "Acceptance Criteria", "Procedure",
    ]
    register.append(register_headers)
    for procedure in procedures:
        register.append([
            procedure["id"], procedure["title"], procedure["category"],
            procedure["status"], procedure["owner"], " ".join(procedure["stages"]),
            " ".join(procedure["requirements"]),
            " ".join(procedure["design_items"]),
            " ".join(procedure["equipment"]), procedure["criteria"],
            procedure["description"],
        ])
    style_header(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([22, 32, 30, 14, 24, 20, 34, 30, 30, 58, 70], 1):
        register.column_dimensions[get_column_letter(index)].width = width
    register.freeze_panes = "D2"
    register.auto_filter.ref = register.dimensions

    execution_headers = [
        "Procedure ID", "Title", "Acceptance Criteria", "Executor", "Date",
        "Configuration / Serial Number", "Result", "Measured Value / Observation",
        "Evidence Reference", "Deviation / Disposition", "Approval",
    ]
    for stage in ("FAT", "SAT", "COMMISSIONING"):
        selected = [item for item in procedures if stage in item["stages"]]
        if not selected:
            continue
        sheet = workbook.create_sheet(stage)
        sheet.append(execution_headers)
        for procedure in selected:
            sheet.append([
                procedure["id"], procedure["title"], procedure["criteria"],
                None, None, None, None, None, None, None, None,
            ])
        style_header(sheet)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row[0].row].height = 72
        result_validation = DataValidation(
            type="list", formula1='"PASS,FAIL,N/A,NOT RUN"', allow_blank=True
        )
        result_validation.error = "Select a result from the list."
        result_validation.showErrorMessage = True
        sheet.add_data_validation(result_validation)
        result_validation.add(f"G2:G{sheet.max_row}")
        for index, width in enumerate([22, 32, 55, 24, 16, 30, 16, 48, 32, 48, 24], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions

    apply_workbook_metadata(workbook, metadata)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    if not args.validate and args.output is None:
        parser.error("output is required unless --validate is used")
    try:
        metadata, metadata_warnings = load_and_validate_metadata(args.input)
        procedures = parse_procedures(args.input)
        warnings = validate_procedures(procedures)
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    warnings = [f"metadata: {warning}" for warning in metadata_warnings] + warnings
    for warning in warnings:
        print(f"WARNING: {warning}")
    if args.validate:
        print(
            f"Validation passed: {len(procedures)} procedures, "
            f"{len(warnings)} warning(s)"
        )
    else:
        export_xlsx(procedures, args.output, metadata)
        print(f"Exported {len(procedures)} procedures to {args.output}")


if __name__ == "__main__":
    main()
