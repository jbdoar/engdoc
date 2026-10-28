"""Tests for the shared Org reader and Python exporters."""

from datetime import date
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [
    str(ROOT / "documents"),
    str(ROOT / "documents" / "project"),
    str(ROOT / "documents" / "requirements"),
    str(ROOT / "documents" / "design"),
    str(ROOT / "documents" / "acceptance"),
    str(ROOT / "documents" / "calibration"),
    str(ROOT / "documents" / "risk"),
    str(ROOT / "documents" / "installation"),
    str(ROOT / "documents" / "service-manual"),
    str(ROOT / "documents" / "user-manual"),
]

from engdoc_org import read_headings
from export_bom import parse_bom, write_xlsx as write_bom, export_bom
from export_schedule import parse_project, write_gantt
from export_requirements import parse_requirements, validate_requirements, export_xlsx
from export_checklists import export_checklist
from validate_project import validate_project
from engdoc_metadata import load_and_validate_metadata
from export_traceability import validate_project_set, export_traceability
from export_design import (
    parse_design_items,
    validate_design,
    export_xlsx as export_design_xlsx,
)
from export_acceptance import (
    parse_procedures,
    validate_procedures,
    export_xlsx as export_acceptance_xlsx,
)
from export_calibration import (
    parse_equipment,
    validate_equipment,
    export_xlsx as export_calibration_xlsx,
)
from export_risk import parse_hazards, validate_hazards, export_xlsx as export_risk_xlsx
from export_installation import (
    parse_activities, validate_activities, export_xlsx as export_installation_xlsx,
)
from export_service import parse_tasks, validate_tasks, export_xlsx as export_service_xlsx
from export_user_manual import (
    parse_topics, validate_topics, export_xlsx as export_user_xlsx,
)


class OrgReaderTests(unittest.TestCase):
    def test_heading_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.org"
            source.write_text(
                "* Parent\n** TODO Child :BOM:@alex:\n"
                ":PROPERTIES:\n:PN: 0012\n:END:\n"
                "SCHEDULED: <2026-10-27 Tue>\nBody\n",
                encoding="utf-8",
            )
            heading = read_headings(source, ["TODO"])[1]
            self.assertEqual(heading.path, ["Parent"])
            self.assertEqual(heading.title, "Child")
            self.assertEqual(heading.todo, "TODO")
            self.assertEqual(heading.tags, ["BOM", "@alex"])
            self.assertEqual(heading.properties["PN"], "0012")
            self.assertIn("SCHEDULED", heading.planning)


class ExporterTests(unittest.TestCase):
    def test_bom_external_format_adapter(self):
        rows = parse_bom(ROOT / "test" / "project.org")
        metadata, _warnings = load_and_validate_metadata(ROOT / "test" / "project.org")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            adapter_dir = root / "bom" / "example-erp"
            adapter_dir.mkdir(parents=True)
            (adapter_dir / "adapter.py").write_text(
                "from openpyxl import Workbook\n"
                "def export(rows, output, metadata):\n"
                "    wb = Workbook()\n"
                "    ws = wb.active\n"
                "    ws.title = 'ERP Import'\n"
                "    ws.append(['Part Number', 'Quantity'])\n"
                "    for row in rows:\n"
                "        ws.append([row.get('PN', ''), row.get('QTY', '')])\n"
                "    wb.save(output)\n",
                encoding="utf-8",
            )
            output = root / "erp.xlsx"
            export_bom(rows, output, metadata, "example-erp", [root])
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["ERP Import"])
            self.assertEqual(workbook["ERP Import"]["A1"].value, "Part Number")

    def test_bom_rejects_unknown_format_adapter(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "Unknown BOM format"):
                export_bom([], Path(directory) / "bom.xlsx", {}, "missing", [])

    def test_user_manual_template_and_register(self):
        source = ROOT / "documents" / "user-manual" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        topics = parse_topics(source)
        self.assertEqual(validate_topics(topics), [])
        self.assertEqual(len(topics), 4)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "user-manual.xlsx"
            export_user_xlsx(topics, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["User Content Register", "Metadata"])
            sheet = workbook["User Content Register"]
            self.assertEqual(sheet["A2"].value, "USR-SAFE-001")
            self.assertEqual(sheet.freeze_panes, "G2")

    def test_user_manual_rejects_invalid_audience(self):
        source = ROOT / "test" / "user-manual.org"
        text = source.read_text(encoding="utf-8")
        with tempfile.TemporaryDirectory() as directory:
            invalid = Path(directory) / "user-manual.org"
            invalid.write_text(
                text.replace(":AUDIENCE: OPERATOR", ":AUDIENCE: EVERYONE", 1),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "AUDIENCE"):
                validate_topics(parse_topics(invalid))

    def test_service_template_and_blank_records(self):
        source = ROOT / "documents" / "service-manual" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        tasks = parse_tasks(source)
        self.assertEqual(validate_tasks(tasks), [])
        self.assertEqual(len(tasks), 3)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "service.xlsx"
            export_service_xlsx(tasks, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(
                workbook.sheetnames,
                ["READ ME", "Maintenance Register", "Maintenance Schedule",
                 "Service Record", "Metadata"],
            )
            self.assertIn("NOT A COMPLETED RECORD", workbook["READ ME"]["A1"].value)
            self.assertEqual(workbook["Maintenance Schedule"].max_row, 3)
            record = workbook["Service Record"]
            self.assertIsNone(record["C2"].value)
            self.assertIsNone(record["F2"].value)

    def test_service_rejects_missing_interval_and_invalid_type(self):
        source = ROOT / "test" / "service.org"
        text = source.read_text(encoding="utf-8")
        with tempfile.TemporaryDirectory() as directory:
            invalid = Path(directory) / "service.org"
            invalid.write_text(text.replace(":INTERVAL_DAYS: 180", ":INTERVAL_DAYS:", 1),
                               encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "requires INTERVAL_DAYS"):
                validate_tasks(parse_tasks(invalid))
            invalid.write_text(text.replace(":TASK_TYPE: INSPECTION",
                                            ":TASK_TYPE: UNKNOWN", 1),
                               encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "TASK_TYPE"):
                validate_tasks(parse_tasks(invalid))

    def test_installation_template_and_blank_site_workbook(self):
        source = ROOT / "documents" / "installation" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        activities = parse_activities(source)
        self.assertEqual(validate_activities(activities), [])
        self.assertEqual(len(activities), 4)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "installation.xlsx"
            export_installation_xlsx(activities, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(
                workbook.sheetnames,
                ["READ ME", "Activity Register", "Site Execution", "Punch List", "Metadata"],
            )
            self.assertIn("NOT A COMPLETED RECORD", workbook["READ ME"]["A1"].value)
            execution = workbook["Site Execution"]
            self.assertIsNone(execution["E2"].value)
            self.assertIsNone(execution["H2"].value)
            self.assertEqual(len(execution.data_validations.dataValidation), 1)
            self.assertEqual(workbook["Punch List"].max_row, 21)

    def test_installation_rejects_invalid_phase_and_missing_criteria(self):
        source = ROOT / "test" / "installation.org"
        text = source.read_text(encoding="utf-8")
        with tempfile.TemporaryDirectory() as directory:
            invalid = Path(directory) / "installation.org"
            invalid.write_text(text.replace(":PHASE: SITE", ":PHASE: FIELD", 1),
                               encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "PHASE"):
                validate_activities(parse_activities(invalid))
            invalid.write_text(
                text.replace(":ACCEPTANCE_CRITERIA: Site access, utilities, environment, and interfaces meet approved criteria.\n", "", 1),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "ACCEPTANCE_CRITERIA"):
                validate_activities(parse_activities(invalid))

    def test_risk_template_and_register(self):
        source = ROOT / "documents" / "risk" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        hazards = parse_hazards(source)
        self.assertEqual(validate_hazards(hazards), [])
        self.assertEqual(hazards[0]["initial_risk"], 12)
        self.assertEqual(hazards[0]["residual_risk"], 4)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "risk.xlsx"
            export_risk_xlsx(hazards, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames,
                             ["Hazard Register", "Risk Matrix", "Metadata"])
            self.assertEqual(workbook["Hazard Register"]["A2"].value,
                             "HAZ-MECH-001")
            self.assertEqual(workbook["Risk Matrix"]["F2"].value, 25)

    def test_risk_rejects_invalid_scores_and_residual_increase(self):
        source = ROOT / "test" / "risk.org"
        hazards = parse_hazards(source)
        hazards[0]["residual_risk"] = 20
        with self.assertRaisesRegex(ValueError, "residual risk exceeds"):
            validate_hazards(hazards)

        text = source.read_text(encoding="utf-8").replace(":SEVERITY: 4", ":SEVERITY: 6", 1)
        with tempfile.TemporaryDirectory() as directory:
            invalid = Path(directory) / "risk.org"
            invalid.write_text(text, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "1 to 5"):
                parse_hazards(invalid)

    def test_calibration_template_and_blank_execution_workbook(self):
        source = ROOT / "documents" / "calibration" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        equipment = parse_equipment(source)
        self.assertEqual(validate_equipment(equipment, date(2026, 10, 27)), [])
        self.assertEqual(len(equipment), 3)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "calibration.xlsx"
            export_calibration_xlsx(equipment, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(
                workbook.sheetnames,
                ["READ ME", "Equipment Register", "Calibration Execution", "Metadata"],
            )
            self.assertIn("NOT A COMPLETED RECORD", workbook["READ ME"]["A1"].value)
            execution = workbook["Calibration Execution"]
            self.assertEqual(execution["A2"].value, "EQ-STD-001")
            self.assertIsNone(execution["D2"].value)
            self.assertIsNone(execution["K2"].value)
            self.assertEqual(len(execution.data_validations.dataValidation), 1)

    def test_calibration_rejects_bad_standards_and_reports_overdue(self):
        source = ROOT / "test" / "calibration.org"
        equipment = parse_equipment(source)
        equipment[1]["due"] = "2020-01-01"
        warnings = validate_equipment(equipment, date(2026, 10, 27))
        self.assertIn("overdue", warnings[0])

        equipment = parse_equipment(source)
        equipment[0]["standards"] = ["EQ-DMM-001"]
        with self.assertRaisesRegex(ValueError, "cycle"):
            validate_equipment(equipment, date(2026, 10, 27))

        equipment = parse_equipment(source)
        equipment[1]["standards"] = ["EQ-MISSING"]
        with self.assertRaisesRegex(ValueError, "unknown reference standard"):
            validate_equipment(equipment, date(2026, 10, 27))

    def test_acceptance_template_and_blank_execution_workbook(self):
        source = ROOT / "documents" / "acceptance" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        procedures = parse_procedures(source)
        self.assertEqual(validate_procedures(procedures), [])
        self.assertEqual(len(procedures), 3)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "acceptance.xlsx"
            export_acceptance_xlsx(procedures, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(
                workbook.sheetnames,
                ["READ ME", "Procedures", "FAT", "SAT", "Metadata"],
            )
            self.assertIn("NOT A COMPLETED RECORD", workbook["READ ME"]["A1"].value)
            fat = workbook["FAT"]
            self.assertEqual(fat["A2"].value, "ATP-SYS-001")
            self.assertIsNone(fat["D2"].value)
            self.assertIsNone(fat["G2"].value)
            self.assertEqual(len(fat.data_validations.dataValidation), 1)
            self.assertEqual(workbook.properties.identifier, "ENG-0001-ATP")

    def test_acceptance_rejects_invalid_procedures(self):
        base = (
            "#+TITLE: Acceptance\n#+AUTHOR: Engineer\n"
            "#+ENGDOC_PROJECT_NAME: Project\n#+ENGDOC_PROJECT_ID: P-1\n"
            "#+ENGDOC_DOCUMENT_ID: P-1-ATP\n#+ENGDOC_REVISION: A\n"
            "#+ENGDOC_STATUS: DRAFT\n"
        )
        cases = {
            "status": (":STATUS: INVALID\n:OWNER: Test\n:STAGES: FAT\n"
                       ":ACCEPTANCE_CRITERIA: Pass", "STATUS"),
            "stage": (":STATUS: DRAFT\n:OWNER: Test\n:STAGES: FIELD\n"
                      ":ACCEPTANCE_CRITERIA: Pass", "invalid stage"),
            "owner": (":STATUS: DRAFT\n:STAGES: FAT\n"
                      ":ACCEPTANCE_CRITERIA: Pass", "missing OWNER"),
            "criteria": (":STATUS: DRAFT\n:OWNER: Test\n:STAGES: FAT",
                         "missing ACCEPTANCE_CRITERIA"),
        }
        for name, (properties, message) in cases.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "acceptance.org"
                source.write_text(
                    base + "* Tests\n** Procedure\n:PROPERTIES:\n"
                    ":CUSTOM_ID: ATP-TEST-001\n" + properties
                    + "\n:END:\nExecute the procedure.\n",
                    encoding="utf-8",
                )
                with self.assertRaisesRegex(ValueError, message):
                    procedures = parse_procedures(source)
                    validate_procedures(procedures)

    def test_design_template_and_register(self):
        source = ROOT / "documents" / "design" / "template.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        items = parse_design_items(source)
        self.assertEqual(validate_design(items), [])
        self.assertEqual(len(items), 3)
        self.assertEqual(items[0]["id"], "DES-SYS-001")

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "design.xlsx"
            export_design_xlsx(items, output, metadata)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["Design Register", "Metadata"])
            sheet = workbook["Design Register"]
            self.assertEqual(sheet["A2"].value, "DES-SYS-001")
            self.assertEqual(sheet["E2"].value, "DRAFT")
            self.assertEqual(sheet.freeze_panes, "D2")
            self.assertEqual(workbook.properties.identifier, "ENG-0001-DESIGN")

    def test_design_rejects_invalid_items(self):
        base = (
            "#+TITLE: Design\n#+AUTHOR: Engineer\n"
            "#+ENGDOC_PROJECT_NAME: Project\n#+ENGDOC_PROJECT_ID: P-1\n"
            "#+ENGDOC_DOCUMENT_ID: P-1-DESIGN\n#+ENGDOC_REVISION: A\n"
            "#+ENGDOC_STATUS: DRAFT\n"
        )
        cases = {
            "status": (":STATUS: UNKNOWN\n:OWNER: Engineer", "invalid STATUS"),
            "owner": (":STATUS: DRAFT", "missing OWNER"),
            "description": (":STATUS: DRAFT\n:OWNER: Engineer", "missing design description"),
        }
        for name, (properties, message) in cases.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "design.org"
                description = "\nDesign rationale.\n" if name != "description" else "\n"
                source.write_text(
                    base + "* Category\n** Item\n:PROPERTIES:\n"
                    ":CUSTOM_ID: DES-TEST-001\n" + properties + "\n:END:\n" + description,
                    encoding="utf-8",
                )
                with self.assertRaisesRegex(ValueError, message):
                    items = parse_design_items(source)
                    validate_design(items)

    def test_project_set_traceability(self):
        model = validate_project_set(ROOT / "test")
        self.assertEqual(len(model["documents"]), 9)
        self.assertEqual(
            {link["heading"] for link in model["links"]["REQ-SYS-001"]},
            {"Assembly", "Control architecture", "Startup sequence test",
             "Unexpected motion", "Execute controlled startup",
             "Inspect protective functions", "Safety guidance"},
        )
        self.assertEqual(model["links"]["REQ-SYS-002"], [])

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "traceability.xlsx"
            export_traceability(model, output)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames,
                             ["Traceability", "Documents", "Design Coverage",
                              "Equipment Coverage", "Hazard Coverage",
                              "Installation Coverage", "Service Coverage",
                              "User Guidance Coverage", "Metadata"])
            sheet = workbook["Traceability"]
            self.assertEqual(sheet["A2"].value, "REQ-SYS-001")
            self.assertIn("DES-SYS-001", sheet["G2"].value)
            self.assertEqual(sheet["I2"].value, "LINKED")
            requirement_rows = {
                sheet.cell(row, 1).value: sheet.cell(row, 9).value
                for row in range(2, sheet.max_row + 1)
            }
            self.assertEqual(requirement_rows["REQ-SYS-002"], "UNLINKED")
            self.assertEqual(workbook.properties.identifier, "ENG-0001-TRACE")
            design = workbook["Design Coverage"]
            self.assertEqual(design["A2"].value, "DES-SYS-001")
            self.assertIn("ATP-SYS-001", design["E2"].value)
            self.assertEqual(design["G2"].value, "COVERED")
            equipment = workbook["Equipment Coverage"]
            equipment_rows = {
                equipment.cell(row, 1).value: equipment.cell(row, 7).value
                for row in range(2, equipment.max_row + 1)
            }
            self.assertEqual(equipment_rows["EQ-DMM-001"], "ASSIGNED")
            self.assertEqual(equipment_rows["EQ-STD-001"], "REFERENCE STANDARD")
            hazards = workbook["Hazard Coverage"]
            self.assertEqual(hazards["A2"].value, "HAZ-MECH-001")
            self.assertEqual(hazards["I2"].value, "CONTROLLED")
            installation = workbook["Installation Coverage"]
            self.assertEqual(installation["A2"].value, "INS-SITE-001")
            self.assertEqual(installation["H2"].value, "CONTROLLED")
            service = workbook["Service Coverage"]
            self.assertEqual(service["A2"].value, "SVC-SAFE-001")
            self.assertEqual(service["G2"].value, "CONTROLLED")
            user = workbook["User Guidance Coverage"]
            self.assertEqual(user["A2"].value, "USR-SAFE-001")
            self.assertEqual(user["G2"].value, "CONTROLLED")

    def test_project_set_rejects_broken_links(self):
        cases = {
            "requirement": ("project.org", "REQ-SYS-001", "REQ-NOPE", "unknown requirement"),
            "design-item": (
                "acceptance.org", "DES-SYS-001", "DES-NOPE", "unknown design item",
            ),
            "equipment": (
                "acceptance.org", "EQ-DMM-001", "EQ-NOPE", "unknown equipment",
            ),
            "acceptance-procedure": (
                "risk.org", "ATP-SYS-001", "ATP-NOPE",
                "unknown acceptance procedure",
            ),
            "hazard": (
                "installation.org", "HAZ-PERF-001", "HAZ-NOPE", "unknown hazard",
            ),
            "service-hazard": (
                "service.org", "HAZ-MECH-001", "HAZ-NOPE", "unknown hazard",
            ),
            "service-task": (
                "user-manual.org", "SVC-SAFE-001", "SVC-NOPE",
                "unknown service task",
            ),
            "document": (
                "project.org", "ENG-0001-REQ", "ENG-0001-MISSING",
                "unknown related document",
            ),
            "project": (
                "requirements.org", "ENGDOC_PROJECT_ID: ENG-0001",
                "ENGDOC_PROJECT_ID: OTHER", "different project IDs",
            ),
        }
        for name, (filename, old, new, message) in cases.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                target = Path(directory)
                shutil.copy(ROOT / "test" / "project.org", target / "project.org")
                shutil.copy(ROOT / "test" / "requirements.org", target / "requirements.org")
                shutil.copy(ROOT / "test" / "design.org", target / "design.org")
                shutil.copy(ROOT / "test" / "acceptance.org", target / "acceptance.org")
                shutil.copy(ROOT / "test" / "calibration.org", target / "calibration.org")
                shutil.copy(ROOT / "test" / "risk.org", target / "risk.org")
                shutil.copy(ROOT / "test" / "installation.org", target / "installation.org")
                shutil.copy(ROOT / "test" / "service.org", target / "service.org")
                shutil.copy(ROOT / "test" / "user-manual.org",
                            target / "user-manual.org")
                path = target / filename
                path.write_text(
                    path.read_text(encoding="utf-8").replace(old, new, 1),
                    encoding="utf-8",
                )
                with self.assertRaisesRegex(ValueError, message):
                    validate_project_set(target)

    def test_project_parsers(self):
        bom = parse_bom(ROOT / "test" / "project.org")
        schedule = parse_project(ROOT / "test" / "project.org")
        self.assertEqual([row["ITEM"] for row in bom[:2]],
                         ["System Assembly", "Component or Subassembly"])
        self.assertTrue(any(row["TODO"] == "TODO" for row in schedule))

    def test_schedule_dates_and_owner(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "project.org"
            source.write_text(
                "* Work\n** NEXT Build :@sam:\n"
                "SCHEDULED: <2026-10-27 Tue> DEADLINE: <2026-10-29 Thu>\n",
                encoding="utf-8",
            )
            task = parse_project(source)[1]
            self.assertEqual(task["OWNERS"], ["sam"])
            self.assertEqual(task["START"], date(2026, 10, 27))
            self.assertEqual(task["END"], date(2026, 10, 29))

    def test_requirements_and_workbooks(self):
        source = ROOT / "test" / "requirements.org"
        metadata, metadata_warnings = load_and_validate_metadata(source)
        self.assertEqual(len(metadata_warnings), 1)
        requirements = parse_requirements(source)
        self.assertEqual(validate_requirements(requirements), 1)
        self.assertEqual(len({item["id"] for item in requirements}), len(requirements))

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            project_source = ROOT / "test" / "project.org"
            project_metadata, _warnings = load_and_validate_metadata(project_source)
            write_bom(parse_bom(project_source), output / "bom.xlsx", project_metadata)
            write_gantt(parse_project(project_source), output / "schedule.xlsx",
                        project_metadata)
            export_xlsx(requirements, output / "requirements.xlsx", metadata)
            export_checklist(requirements, "FAT", output / "fat.xlsx", metadata)
            export_checklist(requirements, "SAT", output / "sat.xlsx", metadata)
            for artifact in output.iterdir():
                self.assertGreater(artifact.stat().st_size, 0)

            bom_book = load_workbook(output / "bom.xlsx")
            bom_sheet = bom_book["BOM"]
            self.assertEqual(
                [cell.value for cell in bom_sheet[1]][:5],
                ["LEVEL", "ITEM", "STATE", "QTY", "PN"],
            )
            self.assertEqual(bom_sheet["B2"].value, "System Assembly")
            self.assertIsInstance(bom_sheet["D2"].value, int)
            self.assertEqual(bom_sheet.row_dimensions[3].outlineLevel, 2)
            self.assertEqual(bom_sheet.freeze_panes, "A2")
            self.assertIn("Metadata", bom_book.sheetnames)
            self.assertEqual(bom_book.properties.identifier, "ENG-0001-PROJECT")

            schedule_book = load_workbook(output / "schedule.xlsx")
            schedule_sheet = schedule_book["Schedule"]
            self.assertEqual(
                [cell.value for cell in schedule_sheet[1]][:5],
                ["Task", "Owner", "Status", "Start", "End"],
            )
            self.assertEqual(schedule_sheet.freeze_panes, "F2")

            requirements_book = load_workbook(output / "requirements.xlsx")
            requirements_sheet = requirements_book["Requirements"]
            self.assertEqual(requirements_sheet.freeze_panes, "B2")
            self.assertEqual(requirements_sheet.auto_filter.ref,
                             requirements_sheet.dimensions)
            self.assertEqual(requirements_sheet["A2"].value, "REQ-SYS-001")
            self.assertEqual(requirements_book["Metadata"]["B4"].value,
                             "ENG-0001")
            self.assertEqual(requirements_book.properties.identifier,
                             "ENG-0001-REQ")

            fat_book = load_workbook(output / "fat.xlsx")
            fat_sheet = fat_book["FAT"]
            self.assertEqual(fat_sheet.freeze_panes, "C2")
            self.assertEqual(fat_sheet["F2"].value, None)
            self.assertEqual(len(fat_sheet.data_validations.dataValidation), 1)
            validation = fat_sheet.data_validations.dataValidation[0]
            self.assertEqual(validation.formula1, '"PASS,FAIL,N/A,NOT RUN"')
            self.assertIn("F2", str(validation.sqref))

    def test_template_is_valid(self):
        requirements = parse_requirements(
            ROOT / "documents" / "requirements" / "template.org"
        )
        self.assertEqual(validate_requirements(requirements), 0)

    def test_project_template_is_valid(self):
        warnings = validate_project(ROOT / "documents" / "project" / "template.org")
        self.assertTrue(any("Specify component" in warning for warning in warnings))

    def test_project_rejects_invalid_bom_and_schedule(self):
        cases = {
            "quantity": (":QTY: 0", "positive integer"),
            "cost": (":QTY: 1\n:COST: nope", "nonnegative cost"),
            "source": (":QTY: 1\n:SOURCE: BORROW", "invalid SOURCE"),
            "dates": (
                ":QTY: 1\n:END:\nSCHEDULED: <2026-10-30 Fri> "
                "DEADLINE: <2026-10-29 Thu>",
                "precedes",
            ),
        }
        for name, (content, message) in cases.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "project.org"
                if ":END:" not in content:
                    content += "\n:END:"
                source.write_text(
                    "#+TITLE: Test Project\n#+AUTHOR: Engineer\n"
                    "#+ENGDOC_PROJECT_NAME: Test Project\n"
                    "#+ENGDOC_PROJECT_ID: P-1\n"
                    "#+ENGDOC_DOCUMENT_ID: P-1-PROJECT\n"
                    "#+ENGDOC_REVISION: A\n#+ENGDOC_STATUS: DRAFT\n"
                    f"* TODO Item :BOM:\n:PROPERTIES:\n{content}\n",
                    encoding="utf-8",
                )
                with self.assertRaisesRegex(ValueError, message):
                    validate_project(source)

    def test_requirements_reject_invalid_input(self):
        invalid = {
            "duplicate": (
                "** A\n:PROPERTIES:\n:CUSTOM_ID: REQ-A\n:VERIFY_STAGE: FAT\n:END:\nA\n"
                "** B\n:PROPERTIES:\n:CUSTOM_ID: REQ-A\n:VERIFY_STAGE: SAT\n:END:\nB\n",
                "Duplicate",
            ),
            "stage": (
                "** A\n:PROPERTIES:\n:CUSTOM_ID: REQ-A\n:VERIFY_STAGE: FIELD\n:END:\nA\n",
                "Unrecognized",
            ),
        }
        for name, (content, message) in invalid.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "requirements.org"
                source.write_text(content, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, message):
                    requirements = parse_requirements(source)
                    validate_requirements(requirements)

    def test_metadata_rejects_missing_invalid_and_unreleased_dates(self):
        base = (
            "#+TITLE: Test\n#+AUTHOR: Engineer\n"
            "#+ENGDOC_PROJECT_NAME: Project\n#+ENGDOC_PROJECT_ID: P-1\n"
            "#+ENGDOC_DOCUMENT_ID: P-1-TEST\n#+ENGDOC_REVISION: A\n"
        )
        cases = {
            "missing": ("#+ENGDOC_STATUS: DRAFT\n", "Missing document metadata"),
            "status": (
                "#+ENGDOC_STATUS: APPROVED\n#+ENGDOC_ISSUE_DATE: 2026-10-27\n",
                "Invalid Status",
            ),
            "date": (
                "#+ENGDOC_STATUS: RELEASED\n#+ENGDOC_ISSUE_DATE: 10/27/2026\n",
                "YYYY-MM-DD",
            ),
            "released-date": ("#+ENGDOC_STATUS: RELEASED\n", "require an Issue Date"),
        }
        for name, (suffix, message) in cases.items():
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "document.org"
                content = suffix if name == "missing" else base + suffix
                source.write_text(content, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, message):
                    load_and_validate_metadata(source)


if __name__ == "__main__":
    unittest.main()
